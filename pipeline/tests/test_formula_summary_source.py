from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook

TEMPLATE = Path(__file__).resolve().parents[1] / "templates" / "dataos_harbor"
if str(TEMPLATE) not in sys.path:
    sys.path.insert(0, str(TEMPLATE))

from __verifier_lib.sources.xlsx import summarize_formulas  # noqa: E402


def test_formula_summary_is_layout_independent(tmp_path: Path) -> None:
    path = tmp_path / "model.xlsx"
    workbook = Workbook()
    inputs = workbook.active
    inputs.title = "Inputs chosen by author"
    inputs["A1"] = 10
    calculations = workbook.create_sheet("Any calculation layout")
    calculations["B7"] = "='Inputs chosen by author'!A1*2"
    calculations["D11"] = "=1+1"
    summary = workbook.create_sheet("Summary with arbitrary name")
    summary["Z4"] = "='Any calculation layout'!B7"
    workbook.save(path)

    result = summarize_formulas(path)

    assert result.n_formula == 3
    assert result.n_cross_sheet_formula == 2
    assert result.n_formula_sheets == 2
    assert result.formula_sheets == [
        "Any calculation layout",
        "Summary with arbitrary name",
    ]
