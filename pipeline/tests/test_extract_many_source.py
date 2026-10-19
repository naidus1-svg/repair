from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from openpyxl import Workbook

TEMPLATE = Path(__file__).resolve().parents[1] / "templates" / "dataos_harbor"
if str(TEMPLATE) not in sys.path:
    sys.path.insert(0, str(TEMPLATE))

from __verifier_lib.models import Source  # noqa: E402
from __verifier_lib.sources.registry import SourceRegistry  # noqa: E402


def test_extract_many_combines_labeled_heterogeneous_files(
    tmp_path: Path,
) -> None:
    workspace = tmp_path / "workspace"
    (workspace / "data").mkdir(parents=True)
    (workspace / "output").mkdir()

    document = Document()
    document.add_paragraph("authoritative source statement")
    document.save(workspace / "data" / "source.docx")

    workbook = Workbook()
    workbook.active["A1"] = "deliverable conclusion"
    workbook.save(workspace / "output" / "model.xlsx")

    source = Source.model_validate(
        {
            "type": "file",
            "file": {
                "type": "filesystem",
                "command": "extract_many",
                "arguments": {
                    "paths": ["data/source.docx", "output/model.xlsx"]
                },
            },
        }
    )
    result = SourceRegistry(workspace, max_content_chars=2000).run(source)

    assert result.paths == ["data/source.docx", "output/model.xlsx"]
    assert "[File: data/source.docx]" in result.text
    assert "authoritative source statement" in result.text
    assert "[File: output/model.xlsx]" in result.text
    assert "deliverable conclusion" in result.text
