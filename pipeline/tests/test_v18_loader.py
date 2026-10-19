"""M2 tests — task corpus loader + preconditions, exercised on the omega v18 dataset."""

from __future__ import annotations

import os

from pathlib import Path

import pytest

from gdpval_qc.v18_spec import V18Spec
from gdpval_qc.v18_loader import (
    FileContent, Sample, TaskCorpusLoader, _extract_text,
)
from gdpval_qc.precondition import (
    BUILTINS, get_precondition, solution_present,
    solution_has_charts, task_has_verifier,
    task_has_rubric_verifier, task_has_deterministic_verifier,
)


REPO_ROOT     = Path(__file__).resolve().parent.parent
V18_SPEC_PATH = REPO_ROOT / "specs" / "gdpval_task_quality_v18.md"
OMEGA_ROOT    = Path(os.environ.get("GDPVAL_TEST_CORPUS") or (Path(__file__).resolve().parent.parent / "data" / "sample_corpus"))

# Tests that need a multi-task corpus (e.g. counts, batch-level checks) are
# auto-skipped against the bundled single-task sample. Point GDPVAL_TEST_CORPUS
# at a full corpus to run them.
_TASK_COUNT     = len(list(OMEGA_ROOT.glob("task_*"))) if OMEGA_ROOT.is_dir() else 0
_FULL_OMEGA     = _TASK_COUNT >= 5


def _spec() -> V18Spec:
    return V18Spec.from_markdown(V18_SPEC_PATH)


def _skip_if_omega_missing() -> None:
    if not OMEGA_ROOT.is_dir():
        pytest.skip(f"omega dataset not present at {OMEGA_ROOT}")


def _skip_if_not_full_omega() -> None:
    if not _FULL_OMEGA:
        pytest.skip(f"needs ≥5 tasks; corpus at {OMEGA_ROOT} has {_TASK_COUNT}")


def _first_task_id() -> str:
    return TaskCorpusLoader(OMEGA_ROOT, _spec()).task_ids()[0]


# Session-scoped fixture: load the 18-task corpus once for all tests in this
# session. Without this, the same docx/xlsx/pdf extraction work would run
# multiple times across tests that need it (cost ≈ 90s per full load).
@pytest.fixture(scope="session")
def omega_samples() -> list[Sample]:
    if not OMEGA_ROOT.is_dir():
        pytest.skip(f"omega dataset not present at {OMEGA_ROOT}")
    loader = TaskCorpusLoader(OMEGA_ROOT, _spec())
    return loader.load_all()


@pytest.fixture(scope="session")
def omega_task_1() -> Sample:
    """Returns the first available task in the corpus. The omega corpus
    starts at task_1, but the bundled sample corpus starts at task_11 —
    so we don't hardcode the index here."""
    if not OMEGA_ROOT.is_dir():
        pytest.skip(f"omega dataset not present at {OMEGA_ROOT}")
    loader = TaskCorpusLoader(OMEGA_ROOT, _spec())
    return loader.load_sample(loader.task_ids()[0])


# ── Corpus enumeration ─────────────────────────────────────────────────────

def test_task_ids_returns_18_sorted_omega_tasks() -> None:
    _skip_if_not_full_omega()
    loader = TaskCorpusLoader(OMEGA_ROOT, _spec())
    ids = loader.task_ids()
    assert len(ids) == 18, f"expected 18 omega tasks, got {len(ids)}"
    assert ids == sorted(ids), "task_ids must be sorted"
    assert all(tid.startswith("task_") for tid in ids)


# ── Field bindings ─────────────────────────────────────────────────────────

def test_load_sample_populates_default_fields(omega_task_1: Sample) -> None:
    s = omega_task_1
    assert s.task_id.startswith("task_")
    assert s.fields["task_id"] == s.task_id
    # task.toml exposes `domain` -> mapped to `category`
    assert s.fields["category"] != ""
    assert isinstance(s.fields["instruction"], str)
    assert len(s.fields["instruction"]) > 0
    assert isinstance(s.fields["input_files"], list)


def test_solution_files_directory_binding_excludes_reasoning_docx(omega_task_1: Sample) -> None:
    sol = omega_task_1.fields["solution_files"]
    assert isinstance(sol, list)
    assert all(isinstance(fc, FileContent) for fc in sol)

    names = {fc.name for fc in sol}
    assert "Reasoning.docx" not in names, "Reasoning.docx must be excluded"
    assert all("Reasoning.docx" not in fc.name for fc in sol)

    # Files alphabetically sorted
    assert names == set(sorted([fc.name for fc in sol]))


def test_solution_reasoning_file_binding_optional(omega_task_1: Sample) -> None:
    reasoning = omega_task_1.fields.get("solution_reasoning")
    # Either populated (string content) or None (file absent)
    assert reasoning is None or isinstance(reasoning, str)
    # If the first task has a Reasoning.docx, the field must be a non-empty string.
    if (OMEGA_ROOT / omega_task_1.task_id / "solution" / "golden_outputs" / "Reasoning.docx").is_file():
        assert isinstance(reasoning, str)
        assert len(reasoning) > 0


def test_verifier_spec_json_binding_loads_when_present(omega_task_1: Sample) -> None:
    vsj = omega_task_1.fields.get("verifier_spec_json")
    assert isinstance(vsj, str)
    assert "verifiers" in vsj


def test_input_files_excludes_dockerfile_and_compose(omega_task_1: Sample) -> None:
    """Dockerfile and docker-compose.yaml are build-time infrastructure and
    MUST NOT appear in input_files (else quality dims grade them as if
    they were agent inputs, producing false positives)."""
    input_files = omega_task_1.fields["input_files"]
    names = {fc.name for fc in input_files}
    assert "Dockerfile" not in names, \
        "Dockerfile must be excluded from input_files"
    assert "docker-compose.yaml" not in names
    assert "docker-compose.yml" not in names


def test_callback_binding_populated_by_builtin_when_omega_has_evals(omega_task_1: Sample) -> None:
    """As of the verifier_assertion_summary callback being auto-registered as
    a built-in (see quality_framework.v18_verifier_summary), this field is
    populated for omega tasks that have `evals/` directories with verifier
    reports. Pre-summary behaviour was None; the new behaviour is a string
    when evals are present, None when they're absent. Either is fine —
    optional bindings must never raise."""
    val = omega_task_1.fields.get("verifier_assertion_summary")
    assert val is None or isinstance(val, str)


def test_callback_binding_runs_when_registered() -> None:
    _skip_if_omega_missing()
    called_with: list[Sample] = []

    def fake_summary(sample: Sample) -> str:
        called_with.append(sample)
        return f"summary for {sample.task_id}"

    first_tid = _first_task_id()
    loader = TaskCorpusLoader(
        OMEGA_ROOT, _spec(),
        callbacks={"build_verifier_assertion_summary": fake_summary},
    )
    s = loader.load_sample(first_tid)
    assert s.fields["verifier_assertion_summary"] == f"summary for {first_tid}"
    assert len(called_with) == 1


# ── Loading the full batch ─────────────────────────────────────────────────

def test_load_all_returns_one_sample_per_task(omega_samples: list[Sample]) -> None:
    _skip_if_not_full_omega()
    assert len(omega_samples) == 18
    seen_ids = {s.task_id for s in omega_samples}
    assert seen_ids == {f"task_{n}" for s in omega_samples for n in [s.task_id.split("_", 1)[1]]}


# ── Preconditions ──────────────────────────────────────────────────────────

def test_solution_present_true_for_omega_tasks(omega_samples: list[Sample]) -> None:
    for s in omega_samples:
        assert solution_present(s) is True, f"expected solution present for {s.task_id}"


def test_task_has_verifier_true_for_omega_tasks(omega_samples: list[Sample]) -> None:
    for s in omega_samples:
        assert task_has_verifier(s) is True, f"expected verifier for {s.task_id}"


def test_task_has_rubric_verifier_true_for_omega(omega_samples: list[Sample]) -> None:
    """All omega verifier.json entries are rubric assertions."""
    _skip_if_not_full_omega()
    rubric_count = sum(1 for s in omega_samples if task_has_rubric_verifier(s))
    assert rubric_count == 18


def test_solution_has_charts_detects_xlsx_with_real_chart(tmp_path: Path) -> None:
    """Synthesize an xlsx that contains an actual chart object and verify
    chart detection fires. (Many omega xlsx files are formula-only models
    with no chart objects, so we can't rely on the corpus for this assertion.)"""
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference

    # Build a task layout with one chart-containing xlsx in golden_outputs
    td = tmp_path / "task_with_chart"
    (td / "solution" / "golden_outputs").mkdir(parents=True)
    (td / "environment").mkdir()
    (td / "instruction.md").write_text("Test", encoding="utf-8")
    (td / "task.toml").write_text('[metadata]\ndomain = "T"\n', encoding="utf-8")

    wb = Workbook()
    ws = wb.active
    ws.append(["Month", "Revenue"])
    for m, r in [("Jan", 10), ("Feb", 12), ("Mar", 15)]:
        ws.append([m, r])
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=4), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=4))
    ws.add_chart(chart, "D2")
    wb.save(td / "solution" / "golden_outputs" / "with_chart.xlsx")

    root = tmp_path / "root"
    root.mkdir()
    td.rename(root / "task_with_chart")

    loader = TaskCorpusLoader(root, _spec())
    s = loader.load_sample("task_with_chart")
    assert solution_has_charts(s) is True


def test_solution_has_charts_runs_on_all_omega_tasks_without_error(
    omega_samples: list[Sample],
) -> None:
    """End-to-end: chart detection runs cleanly across the full corpus.
    No assertion on the result — many omega xlsx files are formula-only
    workbooks with no chart objects, and that's fine."""
    for s in omega_samples:
        result = solution_has_charts(s)
        assert isinstance(result, bool), \
            f"solution_has_charts must return bool, got {type(result)} for {s.task_id}"


def test_solution_has_charts_false_for_pure_text(tmp_path: Path) -> None:
    """Synthesize a task whose solution is only .docx with no images and no
    chart parts — chart detection should return False."""
    # Build minimal fake task layout
    td = tmp_path / "task_fake"
    (td / "solution" / "golden_outputs").mkdir(parents=True)
    # Make a docx with no images using python-docx
    from docx import Document
    doc = Document()
    doc.add_paragraph("This is a plain-text memo with no images and no charts.")
    doc.save(td / "solution" / "golden_outputs" / "memo.docx")
    (td / "instruction.md").write_text("Test instruction", encoding="utf-8")
    (td / "task.toml").write_text(
        '[metadata]\ndomain = "Test"\n', encoding="utf-8"
    )
    (td / "environment").mkdir()

    # Make a containing root with this one task
    root = tmp_path / "root"
    root.mkdir()
    (td).rename(root / "task_fake")

    loader = TaskCorpusLoader(root, _spec())
    s = loader.load_sample("task_fake")
    # solution_files should contain the docx
    assert s.fields["solution_files"]
    assert solution_has_charts(s) is False


def test_get_precondition_lookup() -> None:
    fn = get_precondition("solution_present")
    assert fn is BUILTINS["solution_present"]

    with pytest.raises(KeyError) as exc:
        get_precondition("nonexistent_precondition")
    assert "Unknown precondition" in str(exc.value)


def test_get_precondition_with_extra() -> None:
    custom = lambda s: True  # noqa: E731
    fn = get_precondition("custom_thing", extra={"custom_thing": custom})
    assert fn is custom


# ── Text extraction edge cases ─────────────────────────────────────────────

def test_extract_text_for_docx_returns_paragraphs(tmp_path: Path) -> None:
    from docx import Document
    p = tmp_path / "x.docx"
    doc = Document()
    doc.add_paragraph("hello world")
    doc.add_paragraph("second line")
    doc.save(p)
    text = _extract_text(p)
    assert "hello world" in text
    assert "second line" in text


def test_extract_text_for_unsupported_extension_doesnt_raise(tmp_path: Path) -> None:
    p = tmp_path / "weird.xyz"
    p.write_bytes(b"\x00\x01\x02\x03")
    # Should not raise
    text = _extract_text(p)
    assert isinstance(text, str)


# ── Required-file binding raises on missing path ───────────────────────────

def test_required_file_binding_raises_when_missing(tmp_path: Path) -> None:
    """A non-optional file binding pointing at a missing path raises."""
    # Build a task missing instruction.md? Actually instruction.md is in default
    # fields, not field_bindings. So we need to construct a fake spec with a
    # required file binding.
    # Easiest: synthesize a task without verifier.json, but the canonical v18
    # spec marks verifier_spec_json as optional. So this test must construct a
    # custom spec.
    from gdpval_qc.v18_spec import (
        V18Spec, JudgeModelConfig, FieldBinding, SynthesisConfig,
        CalibrationConfig, ConfidenceDemotion,
    )

    spec = V18Spec(
        name="test_spec",
        version="18.0",
        cache_bust=None,
        models=[JudgeModelConfig(model="x")],
        field_types={},
        field_bindings=[
            FieldBinding(field_name="must_exist", source="file",
                          path="absent.txt", optional=False),
        ],
        synthesis=SynthesisConfig(),
        calibration=CalibrationConfig(),
        severity_scale={0: "a", 1: "b", 2: "c", 3: "d"},
        require_evidence_quote=True,
        confidence_demotion=ConfidenceDemotion(),
        verdict_tiers=["accept"],
        dimensions=[],
    )

    td = tmp_path / "task_x"
    (td / "environment").mkdir(parents=True)
    (td / "instruction.md").write_text("hi", encoding="utf-8")
    (td / "task.toml").write_text("[metadata]\n", encoding="utf-8")

    root = tmp_path / "root"
    root.mkdir()
    td.rename(root / "task_x")

    loader = TaskCorpusLoader(root, spec)
    with pytest.raises(FileNotFoundError):
        loader.load_sample("task_x")
