"""Behaviour tests for the shipped verifier library, run against BOTH template genres.

Covers the defects the 2026-09-18 pipeline audit found: a source-authoring error zeroing the
whole task, the ArrayFormula crash, the formula-authenticity fields (ref-less padding, numbers
stored as text, decoy tabs), uncached-formula detection/recalc passthrough, and extract_many
(pptx, floor-preserving budget redistribution). Plus the template-sync guard.
"""

from __future__ import annotations

import importlib
import json
import subprocess
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.formula import ArrayFormula

REPO = Path(__file__).resolve().parents[2]
GENRES = {
    "dataos_harbor": (REPO / "pipeline" / "templates" / "dataos_harbor", "__verifier_lib"),
    "harbor": (REPO / "pipeline" / "templates" / "harbor", "verifier_lib"),
}


@pytest.fixture(params=sorted(GENRES))
def lib(request):
    root, pkg = GENRES[request.param]
    for name in [m for m in sys.modules if m == pkg or m.startswith(pkg + ".")]:
        del sys.modules[name]
    sys.path.insert(0, str(root))
    try:
        yield importlib.import_module(pkg)
    finally:
        sys.path.remove(str(root))
        for name in [m for m in sys.modules if m == pkg or m.startswith(pkg + ".")]:
            del sys.modules[name]


def _mod(lib, name):
    return importlib.import_module(f"{lib.__name__}.{name}")


def _model_workbook(path: Path) -> None:
    wb = Workbook()
    inputs = wb.active
    inputs.title = "Inputs"
    for i in range(1, 6):
        inputs.cell(i, 1, i * 10)
    calc = wb.create_sheet("Calc")
    for i in range(1, 6):
        calc.cell(i, 1, f"=Inputs!A{i}*2")
    wb.save(path)


def _criterion(name: str, weight: float, command: str, arguments: dict, path: str, expected) -> dict:
    return {
        "name": name,
        "metadata": {"how_justification": "h", "why_justification": "w", "tag": "t", "grading_role": "process", "importance": "secondary", "weight": weight},
        "source": {"type": "file", "file": {"type": "xlsx", "command": command, "arguments": arguments}},
        "assertion": {"type": "deterministic", "expected": expected, "deterministic": {"path": path, "comparison": "greater_than_equal"}},
    }


def test_source_authoring_error_fails_the_criterion_not_the_task(lib, tmp_path: Path) -> None:
    out = tmp_path / "ws" / "output"
    out.mkdir(parents=True)
    _model_workbook(out / "model.xlsx")
    spec = {
        "task_id": "t",
        "config": {"models": ["x/y"], "temperature": 1.0},
        "verifiers": [
            _criterion("v_pass", 0.5, "formula_summary", {"path": "output/model.xlsx"}, "$.n_formula", 1),
            _criterion("v_pinned", 0.3, "formula_authenticity_check", {"path": "output/model.xlsx", "sheet": "No Such Sheet", "cell_range": "A1:A2"}, "$.n_formula", 1),
            _criterion("v_fail", 0.2, "formula_summary", {"path": "output/model.xlsx"}, "$.n_formula", 10_000),
        ],
    }
    spec_path = tmp_path / "verifier.json"
    spec_path.write_text(json.dumps(spec))
    report = _mod(lib, "runner").run_verifier(spec_path, workspace_dir=tmp_path / "ws", verifier_dir=tmp_path / "logs", agent_logs_dir=tmp_path / "agent")
    assert report["reward"] == pytest.approx(0.5)
    assert "error" not in report                          # no fatal error recorded at task level
    assert report["summary"]["errored"] == 1
    rows = {r["name"]: r["result"] for r in report["assertion_results"]}
    assert rows["v_pass"]["success"] is True
    assert rows["v_pinned"]["success"] is False and "source authoring error" in rows["v_pinned"]["error"]
    assert rows["v_fail"]["success"] is False


def test_malformed_spec_is_still_fatal(lib, tmp_path: Path) -> None:
    spec_path = tmp_path / "verifier.json"
    spec_path.write_text(json.dumps({"task_id": "t", "config": {}, "verifiers": [{"name": "broken"}]}))
    with pytest.raises(Exception):
        _mod(lib, "runner").run_verifier(spec_path, workspace_dir=tmp_path, verifier_dir=tmp_path / "logs", agent_logs_dir=tmp_path / "agent")
    assert (tmp_path / "logs" / "reward.txt").read_text().strip() in {"0", "0.0"}


def test_array_formula_does_not_crash_any_formula_command(lib, tmp_path: Path) -> None:
    path = tmp_path / "array.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"], ws["A2"] = 1, 2
    ws["B1"] = ArrayFormula("B1:B2", "=A1:A2*2")
    ws["C1"] = "=A1+A2"
    wb.save(path)
    x = _mod(lib, "sources.xlsx")
    check = x.check_formula_authenticity(path, "S", "A1:C2")
    assert any(c.is_formula and isinstance(c.formula_text, str) for c in check.cells)
    assert x.summarize_formulas(path).n_formula_with_refs >= 2
    assert x.scan_formula_authenticity(path, []).formula_share_all > 0


def test_formula_scan_ignores_padding_counts_text_numbers_and_is_selector_free(lib, tmp_path: Path) -> None:
    x = _mod(lib, "sources.xlsx")
    live = tmp_path / "live.xlsx"
    _model_workbook(live)
    assert x.summarize_formulas(live).n_formula_with_refs == 5
    assert x.scan_formula_authenticity(live, ["no-such-sheet"]).formula_share_all == pytest.approx(0.5)

    padded = tmp_path / "padded.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Model"
    for i in range(1, 11):
        ws.cell(i, 1, f"{i * 1000:,}")   # number stored as text
        ws.cell(i, 2, "$410")
        ws.cell(i, 3, "=1+1")            # ref-less padding
    ws["D1"] = "=A1"
    wb.save(padded)
    assert x.summarize_formulas(padded).n_formula_with_refs == 1
    scan = x.scan_formula_authenticity(padded, ["model"])
    assert scan.n_hardcoded == 20 and scan.formula_share_all == pytest.approx(1 / 21)

    decoy = tmp_path / "decoy.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Model"
    for i in range(1, 6):
        ws.cell(i, 1, i)
        ws.cell(i, 2, f"=A{i}*2")
    answers = wb.create_sheet("Answers")
    for i in range(1, 101):
        answers.cell(i, 1, i * 1.5)
    wb.save(decoy)
    scan = x.scan_formula_authenticity(decoy, ["model"])
    assert scan.formula_share == pytest.approx(0.5)      # selector view is fooled
    assert scan.formula_share_all < 0.06                 # whole-workbook view is not


def test_uncached_detection_and_recalc_passthrough_without_libreoffice(lib, tmp_path: Path, monkeypatch) -> None:
    x = _mod(lib, "sources.xlsx")
    path = tmp_path / "uncached.xlsx"
    _model_workbook(path)                                 # openpyxl writes formulas with no cached values
    assert x._has_uncached_formula(path) is True
    x._RECALC_CACHE.clear()
    monkeypatch.setattr("shutil.which", lambda _name: None)
    assert x._recalculated_copy_if_uncached(path) == path  # no soffice -> original behaviour, no crash
    assert "[Sheet: Calc]" in x.extract_xlsx(path, 5000)


@pytest.mark.skipif(not any(__import__("shutil").which(b) for b in ("soffice", "libreoffice")), reason="LibreOffice not installed")
def test_recalc_fills_uncached_formulas(lib, tmp_path: Path) -> None:
    x = _mod(lib, "sources.xlsx")
    path = tmp_path / "uncached.xlsx"
    _model_workbook(path)
    x._RECALC_CACHE.clear()
    text = x.extract_xlsx(path, 5000)
    assert "20" in text and "100" in text


def _ctx(lib, ws: Path, tmp_path: Path, cap: int):
    return _mod(lib, "source_types").SourceContext(workspace_dir=ws, agent_logs_dir=tmp_path, max_content_chars=cap)


def test_extract_many_redistributes_unused_budget_without_shrinking_anyone(lib, tmp_path: Path) -> None:
    ws = tmp_path / "ws"
    (ws / "output").mkdir(parents=True)
    (ws / "output" / "small.txt").write_text("tiny")
    (ws / "output" / "big.txt").write_text("x" * 5000)
    cmd = _mod(lib, "sources.filesystem").ExtractMany()
    res = cmd.run(cmd.input_model(paths=["output/small.txt", "output/big.txt"]), _ctx(lib, ws, tmp_path, 4000))
    big = res.text.split("[File: output/big.txt]\n", 1)[1]
    assert big.count("x") > 2000                          # more than the 2000-char even split
    assert "tiny" in res.text and "[TRUNCATED:" in big     # still capped, and visibly so


def test_extract_many_reads_pptx_as_slides_not_zip_bytes(lib, tmp_path: Path) -> None:
    pptx = pytest.importorskip("pptx")
    ws = tmp_path / "ws"
    (ws / "output").mkdir(parents=True)
    deck = pptx.Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Quarterly findings headline"
    deck.save(ws / "output" / "deck.pptx")
    cmd = _mod(lib, "sources.filesystem").ExtractMany()
    res = cmd.run(cmd.input_model(paths=["output/deck.pptx"]), _ctx(lib, ws, tmp_path, 4000))
    assert "Quarterly findings headline" in res.text and "PK" not in res.text[:40]


def test_has_chart_counts_chart_parts_and_refuses_entity_declarations(lib, tmp_path: Path) -> None:
    from openpyxl.chart import BarChart, Reference

    ws_dir = tmp_path / "ws"
    (ws_dir / "output").mkdir(parents=True)
    wb = Workbook()
    ws = wb.active
    for i in range(1, 6):
        ws.cell(i, 1, i)
        ws.cell(i, 2, i * 2)
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=1, max_col=2, min_row=1, max_row=5))
    ws.add_chart(chart, "D2")
    wb.save(ws_dir / "output" / "charted.xlsx")
    plain = Workbook()
    plain.save(ws_dir / "output" / "plain.xlsx")

    cmd = next(c for c in _mod(lib, "sources.xlsx").COMMANDS if c.name == "has_chart")
    ctx = _ctx(lib, ws_dir, tmp_path, 4000)
    out = cmd.run(cmd.input_model(path="output/charted.xlsx"), ctx)
    assert out.chart_count == 1 and out.max_series_count == 2
    assert cmd.run(cmd.input_model(path="output/plain.xlsx"), ctx).chart_count == 0


def test_templates_are_in_sync() -> None:
    proc = subprocess.run([sys.executable, str(REPO / "pipeline" / "tools" / "sync_templates.py")], capture_output=True, text=True)
    assert proc.returncode == 0, proc.stdout
