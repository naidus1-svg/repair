#!/usr/bin/env python3
"""Warn when a golden workbook is close to the verifier XLSX extraction cap.

Distinct from ``source_truncation_check.py``, which flags an *exact* hit on
``max_content_chars``. The 2026-09-17 task_146 review found a golden already at
~88% of a 300k cap: a larger but still-compliant plan overflowed, silently
dropping required tabs.

This check estimates extract_text size from cached cell values (no formula
recompute, so it does not import the in-flight xlsx.py) and compares it to the
cap declared in ``tests/verifier_runtime.pyrun``. Warn at 70% usage; fail if
the estimate meets or exceeds the cap.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

CAP_RE = re.compile(
    r'VERIFIER_MAX_CONTENT[^0-9]{0,80}["\']?(\d+)',
    re.IGNORECASE,
)
DEFAULT_CAP = 90000
HEADROOM_RATIO = 0.70


def resolve_package_dir(task_dir: Path) -> Path | None:
    if (task_dir / "tests").is_dir():
        return task_dir
    nested = task_dir / "task_files"
    if nested.is_dir():
        return nested
    revs = sorted(task_dir.glob("rev_*/task_files"))
    return revs[-1] if revs else None


def read_content_cap(package: Path) -> int:
    runtime = package / "tests" / "verifier_runtime.pyrun"
    if runtime.is_file():
        try:
            text = runtime.read_text(encoding="utf-8", errors="replace")
        except OSError:
            text = ""
        match = CAP_RE.search(text)
        if match:
            return int(match.group(1))
    return DEFAULT_CAP


def estimate_extract_chars(path: Path) -> tuple[int, list[str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, data_only=True, read_only=True)
    total = 0
    sheets: list[str] = []
    try:
        for sheet_name in workbook.sheetnames:
            sheets.append(sheet_name)
            total += len(f"[Sheet: {sheet_name}]") + 1
            for row in workbook[sheet_name].iter_rows():
                cells = [str(cell.value) if cell.value is not None else "" for cell in row]
                if any(cell.strip() for cell in cells):
                    total += len(" | ".join(cells)) + 1
    finally:
        workbook.close()
    return total, sheets


def check_task(task_dir: Path) -> dict:
    package = resolve_package_dir(task_dir)
    result: dict = {
        "task_id": task_dir.name,
        "status": "skip",
        "findings": [],
        "n_findings": 0,
        "package": str(package) if package else None,
    }
    if package is None:
        return result
    golden_dir = package / "solution" / "golden_outputs"
    workbooks = sorted(golden_dir.glob("*.xlsx")) + sorted(golden_dir.glob("*.xlsm"))
    if not workbooks:
        return result

    cap = read_content_cap(package)
    findings: list[dict] = []
    worst = "pass"
    for workbook in workbooks:
        try:
            chars, sheets = estimate_extract_chars(workbook)
        except Exception as exc:  # noqa: BLE001 — keep the batch moving
            findings.append(
                {
                    "kind": "xlsx_estimate_error",
                    "severity": "warn",
                    "path": str(workbook),
                    "evidence": str(exc),
                }
            )
            worst = "warn"
            continue
        ratio = chars / cap if cap else 0.0
        row = {
            "path": str(workbook),
            "estimated_chars": chars,
            "cap": cap,
            "ratio": round(ratio, 4),
            "n_sheets": len(sheets),
        }
        if chars >= cap:
            findings.append(
                {
                    "kind": "xlsx_extraction_at_or_over_cap",
                    "severity": "fail",
                    **row,
                    "evidence": (
                        f"{workbook.name} estimated extract size {chars} "
                        f"meets/exceeds cap {cap}; later sheets can go unseen"
                    ),
                }
            )
            worst = "fail"
        elif ratio >= HEADROOM_RATIO:
            findings.append(
                {
                    "kind": "xlsx_extraction_low_headroom",
                    "severity": "warn",
                    **row,
                    "evidence": (
                        f"{workbook.name} uses {ratio:.1%} of the {cap}-char "
                        "extract cap; a larger compliant workbook can overflow"
                    ),
                }
            )
            if worst != "fail":
                worst = "warn"

    result["findings"] = findings
    result["n_findings"] = len(findings)
    result["status"] = worst if findings else "pass"
    return result


def find_task_dirs(root: Path, task_ids=None, id_prefix: str = "") -> list[Path]:
    if task_ids:
        dirs = []
        for tid in task_ids:
            path = root / tid
            if not path.is_dir():
                print(f"WARNING: task dir not found, skipping: {path}", file=sys.stderr)
                continue
            dirs.append(path)
        return dirs
    return sorted(
        path for path in root.iterdir() if path.is_dir() and path.name.startswith(id_prefix)
    )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("task_ids", nargs="*")
    parser.add_argument("--all", action="store_true")
    parser.add_argument("--root", type=Path, required=True)
    parser.add_argument("--id-prefix", default="")
    parser.add_argument("--out", type=Path, default=None)
    args = parser.parse_args(argv)
    if not args.all and not args.task_ids:
        parser.error("pass one or more task ids, or --all")
    task_dirs = find_task_dirs(
        args.root, None if args.all else args.task_ids, id_prefix=args.id_prefix
    )
    results = [check_task(path) for path in task_dirs]
    report = {
        "generated": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "root": str(args.root.resolve()),
        "n_tasks": len(results),
        "results": results,
    }
    text = json.dumps(report, indent=2)
    if args.out:
        args.out.parent.mkdir(parents=True, exist_ok=True)
        args.out.write_text(text, encoding="utf-8")
    else:
        print(text)
    return 0


if __name__ == "__main__":
    sys.exit(main())
