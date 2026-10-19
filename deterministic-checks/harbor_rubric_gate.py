#!/usr/bin/env python3
"""Mechanical gate for the platform's Harbor / Terminal-Bench task rubric.

The DataOS ``QC`` reviewer grades every task against the 15-check Harbor task
rubric (see docs/harbor-task-rubric.md). On 2026-09-16 12/12 EKW tasks failed
it on the same six checks. This gate encodes the mechanically decidable part of
those checks so a package cannot reach upload with a known-failing item.

Checks (status pass / warn / fail, with evidence):
  instruction_concision   roleplay opener; every input file named by absolute
                          path; every graded deliverable named by absolute path;
                          canary line present; length warning.
  no_extraneous_files     every golden file is either copied by solve.sh AND read
                          by a criterion, or it is extraneous; junk patterns.
  solve_sh_scope          solve.sh copies only graded deliverables; LF endings.
  structured_data_schema  a structured deliverable (xlsx/csv/json) needs a schema
                          block (sheet / column contract) in the instruction.
  anti_cheat_robustness   >= 1 non-rubric (deterministic) assertion; every xlsx
                          deliverable with a formula-driven golden carries a
                          formula-based source; goldens with no live formulas
                          are reported for rebuild.
  judge_injection_guard   the shipped judge system prompt marks source content
                          as untrusted.
  test_instruction_alignment  consumes reports/<task>_TRACE.json if present:
                          any criterion without an instruction anchor fails.

Run:  harbor_rubric_gate.py --root <task_dir|batch_dir> [--out report.json]
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

ROLEPLAY_RE = re.compile(r"^\s*(you are|you're|as (a|an|the)\b|act as|imagine (you|that)|assume (you|the role))", re.I)
CANARY_RE = re.compile(r"harbor-canary|BENCHMARK DATA SHOULD NEVER APPEAR|canary GUID", re.I)
SCHEMA_RE = re.compile(r"\b(sheet|tab|worksheet|column|field|header)s?\b", re.I)
JUNK_RE = re.compile(r"(\.round\d*$|\.bak$|^output_variables\.json$|^__pycache__$|^\.DS_Store$|_INCIDENTS|\.pyc$|\.orig$|~$)")
STRUCTURED_EXT = {".xlsx", ".xlsm", ".csv", ".json"}
INSTRUCTION_WORD_WARN = 550
MIN_GOLDEN_FORMULAS = 5


def resolve_package_dir(task_dir: Path) -> Path | None:
    if (task_dir / "tests").is_dir():
        return task_dir
    nested = task_dir / "task_files"
    if nested.is_dir():
        return nested
    revs = sorted(task_dir.glob("rev_*/task_files"))
    return revs[-1] if revs else None


def genre(package: Path) -> str:
    if (package / "tests" / "verifier.json").is_file():
        return "native"
    if list((package / "tests").glob("deliverables_*/judge.toml")):
        return "rewardkit"
    return "unknown"


def load_criteria(package: Path) -> list[dict]:
    vj = package / "tests" / "verifier.json"
    if not vj.is_file():
        return []
    try:
        return json.loads(vj.read_text(encoding="utf-8")).get("verifiers", [])
    except (OSError, json.JSONDecodeError):
        return []


def criterion_sources(criterion: dict) -> list[dict]:
    src = criterion.get("source", {})
    if src.get("type") == "file":
        return [src.get("file", {})]
    if src.get("type") == "files":
        return list(src.get("files", []))
    return [src]


def graded_output_paths(package: Path) -> set[str]:
    """Deliverable basenames read by any criterion (native) or listed in judge.toml (rewardkit)."""
    names: set[str] = set()
    for c in load_criteria(package):
        blob = json.dumps(c.get("source", {}))
        for m in re.finditer(r"output/([^\"'\s,\]]+)", blob):
            names.add(Path(m.group(1)).name)
    inputs = {p.name for p in (package / "environment" / "initial_workspace").rglob("*") if p.is_file()}
    for toml in (package / "tests").glob("deliverables_*/judge.toml"):
        text = toml.read_text(encoding="utf-8", errors="replace")
        m = re.search(r"^files\s*=\s*\[(.*?)\]", text, re.M | re.S)
        if not m:
            continue
        for f in re.findall(r"[\"']([^\"']+)[\"']", m.group(1)):
            # A path under /tmp/ (e.g. /tmp/judge_views/*__cell_view.md) is a grader-generated
            # judge-only artifact (see the cell-view fix in fix-task-from-review/SKILL.md's
            # G6/uncached-formula mechanism) -- rewritten fresh by test.sh every run, never
            # written by the solver, and never something the instruction should name (the
            # solver has no reason to know it exists). It is not a "graded deliverable" for the
            # instruction_concision/no_extraneous_files/solve_sh_scope purposes below.
            if f.startswith("/tmp/"):
                continue
            if Path(f).name not in inputs:
                names.add(Path(f).name)
    return names


def golden_files(package: Path) -> list[Path]:
    out: list[Path] = []
    for sub in ("solution/golden_outputs", "solution/files"):
        d = package / sub
        if d.is_dir():
            out += [p for p in d.rglob("*") if p.is_file()]
    return out


def input_files(package: Path) -> tuple[str, list[Path]]:
    for sub, prefix in (("environment/data", "/workspace/data"), ("environment/initial_workspace", "/app")):
        d = package / sub
        if d.is_dir():
            return prefix, [p for p in d.rglob("*") if p.is_file() and not JUNK_RE.search(p.name)]
    return "/workspace/data", []


def solve_sh_copies(package: Path) -> tuple[str, set[str]]:
    sh = package / "solution" / "solve.sh"
    if not sh.is_file():
        return "", set()
    text = sh.read_bytes().decode("utf-8", errors="replace")
    names: set[str] = set()
    for m in re.finditer(r"[\"']?([A-Za-z0-9_\-. ]+\.(?:xlsx|xlsm|docx|pdf|pptx|csv|json|txt|md))[\"']?", text):
        names.add(m.group(1).strip())
    if re.search(r"golden_outputs/\*|files/\*|cp\s+-r|cp\s+-a", text):
        names.add("*")
    return text, names


def count_formulas(path: Path) -> int | None:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return None
    try:
        wb = load_workbook(path, data_only=False, read_only=True)
    except Exception:  # noqa: BLE001
        return None
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.startswith("="):
                    n += 1
                elif v is not None and type(v).__name__ in ("ArrayFormula", "DataTableFormula"):
                    n += 1
    return n


def find_trace(package: Path, task_name: str) -> Path | None:
    for parent in [package, *package.parents][:4]:
        cand = parent / "reports" / f"{task_name}_TRACE.json"
        if cand.is_file():
            return cand
    return None


def _res(status: str, evidence: list[str]) -> dict:
    return {"status": status, "evidence": evidence}


def check_instruction_concision(package: Path, graded: set[str]) -> dict:
    ins_path = package / "instruction.md"
    if not ins_path.is_file():
        return _res("fail", ["instruction.md missing"])
    ins = ins_path.read_text(encoding="utf-8", errors="replace")
    ev: list[str] = []
    status = "pass"
    if ROLEPLAY_RE.match(ins):
        ev.append(f"roleplay opener: {ins.strip().splitlines()[0][:80]!r}")
        status = "fail"
    prefix, inputs = input_files(package)
    missing_in = [str(p.relative_to(package / ("environment/data" if prefix == "/workspace/data" else "environment/initial_workspace")))
                  for p in inputs if f"{prefix}/{p.relative_to(package / ('environment/data' if prefix == '/workspace/data' else 'environment/initial_workspace'))}" not in ins]
    if missing_in:
        ev.append(f"{len(missing_in)}/{len(inputs)} input files not named by absolute path: {missing_in[:6]}")
        status = "fail"
    out_prefix = "/app" if prefix == "/app" else "/workspace/output"
    missing_out = [n for n in sorted(graded) if f"{out_prefix}/{n}" not in ins]
    if missing_out:
        ev.append(f"graded deliverables not named as absolute path: {missing_out}")
        status = "fail"
    if not CANARY_RE.search(ins):
        ev.append("no canary line")
        status = "fail"
    words = len(ins.split())
    if words > INSTRUCTION_WORD_WARN:
        ev.append(f"instruction is {words} words (> {INSTRUCTION_WORD_WARN})")
        if status == "pass":
            status = "warn"
    return _res(status, ev or [f"ok ({words} words, {len(inputs)} inputs, {len(graded)} deliverables)"])


def check_no_extraneous_files(package: Path, graded: set[str]) -> dict:
    ev: list[str] = []
    _, copies = solve_sh_copies(package)
    for g in golden_files(package):
        if JUNK_RE.search(g.name):
            ev.append(f"junk in golden dir: {g.name}")
            continue
        if g.name not in graded:
            ev.append(f"golden file not read by any criterion: {g.name}")
        elif "*" not in copies and g.name not in copies:
            ev.append(f"graded golden file not copied by solve.sh: {g.name}")
    for p in package.rglob("*"):
        if p.is_file() and JUNK_RE.search(p.name) and "golden_outputs" not in p.parts and "files" not in p.parts:
            ev.append(f"junk file: {p.relative_to(package)}")
        if p.is_dir() and p.name == "__pycache__":
            ev.append(f"junk dir: {p.relative_to(package)}")
    return _res("fail" if ev else "pass", ev or ["ok"])


def check_solve_sh_scope(package: Path, graded: set[str]) -> dict:
    text, copies = solve_sh_copies(package)
    if not text:
        return _res("skip", ["no solve.sh"])
    ev: list[str] = []
    if "\r\n" in text:
        ev.append("CRLF line endings (exit 127 at runtime)")
    if "*" in copies:
        ev.append("wildcard copy of golden dir")
    extra = sorted(n for n in copies if n != "*" and n not in graded and not n.endswith(".sh"))
    if extra:
        ev.append(f"copies ungraded files: {extra}")
    return _res("fail" if ev else "pass", ev or ["ok"])


def check_structured_data_schema(package: Path, graded: set[str]) -> dict:
    structured = sorted(n for n in graded if Path(n).suffix.lower() in STRUCTURED_EXT)
    if not structured:
        return _res("skip", ["no structured deliverable"])
    ins = (package / "instruction.md").read_text(encoding="utf-8", errors="replace") if (package / "instruction.md").is_file() else ""
    hits = len(SCHEMA_RE.findall(ins))
    if hits == 0:
        return _res("fail", [f"structured deliverables {structured} but instruction never names a sheet/tab/column/field"])
    if hits < 3:
        return _res("warn", [f"only {hits} schema term(s) in instruction for {structured}; confirm a real sheet/column contract exists"])
    return _res("pass", [f"{hits} schema terms for {structured}"])


def check_anti_cheat(package: Path, graded: set[str]) -> dict:
    ev: list[str] = []
    status = "pass"
    g = genre(package)
    if g == "native":
        crit = load_criteria(package)
        det = [c["name"] for c in crit if c.get("assertion", {}).get("type") != "rubric"]
        if not det:
            ev.append(f"0 of {len(crit)} assertions are deterministic (all LLM rubric)")
            status = "fail"
        else:
            ev.append(f"deterministic assertions: {det}")
        formula_sources: dict[str, list[str]] = {}
        for c in crit:
            for s in criterion_sources(c):
                cmd = str(s.get("command", ""))
                if "formula" in cmd or "verify_selections" in cmd:
                    p = Path(str(s.get("arguments", {}).get("path", ""))).name
                    formula_sources.setdefault(p, []).append(c["name"])
    else:
        # rewardkit: a deterministic check is a .py beside the judge.toml; credit it to every graded
        # workbook it names (or to all of them when it discovers the workbook by extension).
        formula_sources = {}
        texts = {p: p.read_text(encoding="utf-8", errors="replace") for p in (package / "tests").rglob("*") if p.suffix in (".toml", ".py", ".sh")}
        if not any("formula" in t for t in texts.values()):
            ev.append("rewardkit package: no formula/deterministic check found under tests/")
            status = "fail"
        workbooks = [n for n in graded if n.lower().endswith((".xlsx", ".xlsm"))]
        for path, text in texts.items():
            if path.suffix != ".py" or "formula" not in text:
                continue
            named = [n for n in workbooks if n in text] or workbooks
            for n in named:
                formula_sources.setdefault(n, []).append(str(path.relative_to(package / "tests")))
    for gf in golden_files(package):
        if gf.suffix.lower() not in (".xlsx", ".xlsm") or gf.name not in graded:
            continue
        n = count_formulas(gf)
        if n is None:
            ev.append(f"{gf.name}: could not count formulas")
            continue
        if n < MIN_GOLDEN_FORMULAS:
            ev.append(f"{gf.name}: golden has {n} live formulas -- hardcoded golden, rebuild before a formula check is possible")
            if status == "pass":
                status = "warn"
        elif gf.name not in formula_sources:
            ev.append(f"{gf.name}: golden has {n} formulas but no criterion checks formula authenticity on it")
            status = "fail"
        else:
            ev.append(f"{gf.name}: {n} formulas, checked by {formula_sources[gf.name]}")
    return _res(status, ev or ["ok"])


def check_judge_guard(package: Path) -> dict:
    for lib in ("__verifier_lib", "verifier_lib"):
        j = package / "tests" / lib / "judge.py"
        if j.is_file():
            text = j.read_text(encoding="utf-8", errors="replace")
            if re.search(r"untrusted", text, re.I) and re.search(r"BEGIN SOURCE|<source>|SOURCE DATA", text):
                return _res("pass", ["judge prompt marks source content untrusted and delimits it"])
            return _res("fail", [f"{lib}/judge.py has no untrusted-content guard / delimiters"])
    if genre(package) == "rewardkit":
        return _res("skip", ["rewardkit genre: judge prompt owned by harbor-rewardkit"])
    return _res("skip", ["no judge.py shipped"])


def check_alignment(package: Path, task_name: str) -> dict:
    trace = find_trace(package, task_name)
    if trace is None:
        return _res("warn", ["no reports/<task>_TRACE.json -- criterion traceability not established"])
    try:
        rows = json.loads(trace.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        return _res("fail", [f"trace unreadable: {exc}"])
    # A TRACE.json may carry ONLY reward_signal_gate's top-level `dispositions` map
    # (its own documented format, see that script's docstring) with no `criteria`
    # list at all -- that's a valid file for G1/G3/G4/G5/G7 disposition purposes,
    # just one that hasn't populated per-criterion instruction anchors yet. Treat
    # that shape the same as a missing file for THIS check (warn, don't crash) --
    # iterating rows.get("criteria", rows) previously fell back to the whole dict
    # and crashed on entries's string keys.
    entries = rows.get("criteria") if isinstance(rows, dict) else rows
    if entries is None:
        return _res("warn", ["reports/<task>_TRACE.json has no 'criteria' list -- criterion traceability not established"])
    names = {c["name"] for c in load_criteria(package)}
    traced = {e.get("name") for e in entries}
    ev: list[str] = []
    for e in entries:
        anchor = e.get("anchor") or e.get("instruction_clause")
        if not anchor or str(anchor).strip().lower() in ("none", "n/a", ""):
            ev.append(f"{e.get('name')}: no instruction anchor")
    untraced = sorted(names - traced)
    if untraced:
        ev.append(f"criteria absent from trace: {untraced}")
    uncovered = [r for r in (rows.get("uncovered_requirements", []) if isinstance(rows, dict) else []) if r]
    if uncovered:
        ev.append(f"instruction requirements with no criterion: {uncovered}")
    return _res("fail" if ev else "pass", ev or [f"all {len(names)} criteria anchored; no uncovered requirement"])


def check_task(task_dir: Path) -> dict:
    package = resolve_package_dir(task_dir)
    if package is None:
        return {"status": "skip", "checks": {}, "evidence": ["no package found"]}
    task_name = re.sub(r"_fixed$", "", task_dir.name) if task_dir.name.startswith("task_") else package.parent.name
    graded = graded_output_paths(package)
    checks = {
        "instruction_concision": check_instruction_concision(package, graded),
        "no_extraneous_files": check_no_extraneous_files(package, graded),
        "solve_sh_scope": check_solve_sh_scope(package, graded),
        "structured_data_schema": check_structured_data_schema(package, graded),
        "anti_cheat_robustness": check_anti_cheat(package, graded),
        "judge_injection_guard": check_judge_guard(package),
        "test_instruction_alignment": check_alignment(package, task_name),
    }
    order = {"fail": 3, "warn": 2, "pass": 1, "skip": 0}
    worst = max(checks.values(), key=lambda c: order[c["status"]])["status"]
    return {"status": worst, "package": str(package), "genre": genre(package), "graded": sorted(graded), "checks": checks}


def discover(root: Path) -> list[Path]:
    root = root.resolve()
    if resolve_package_dir(root) is not None and (root / "tests").is_dir():
        return [root]
    if (root / "task_files").is_dir() or list(root.glob("rev_*/task_files")):
        return [root]
    return sorted(p for p in root.iterdir() if p.is_dir() and p.name.startswith("task_"))


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--root", type=Path, required=True)
    ap.add_argument("--out", type=Path)
    ap.add_argument("--task", action="append", dest="task_ids")
    args = ap.parse_args(argv)
    dirs = discover(args.root)
    if args.task_ids:
        dirs = [d for d in dirs if d.name in set(args.task_ids) or d.name.replace("_fixed", "") in set(args.task_ids)]
    rows = {}
    for d in dirs:
        rows[d.name] = check_task(d)
        r = rows[d.name]
        failing = [k for k, v in r.get("checks", {}).items() if v["status"] == "fail"]
        warning = [k for k, v in r.get("checks", {}).items() if v["status"] == "warn"]
        print(f"{d.name:12} {r['status']:5} fail={failing} warn={warning}")
    report = {"generated": datetime.now(timezone.utc).isoformat(), "root": str(args.root.resolve()), "tasks": rows}
    if args.out:
        args.out.parent.mkdir(parents=True, exist_ok=True)
        args.out.write_text(json.dumps(report, indent=2), encoding="utf-8")
        print(f"wrote {args.out}")
    return 1 if any(r["status"] == "fail" for r in rows.values()) else 0


if __name__ == "__main__":
    sys.exit(main())
