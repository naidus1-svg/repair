#!/usr/bin/env python3
"""Unit tests for source_citation_existence_check.py.

Two layers, matching content_audit_check_test.py's style (self-contained,
no external test framework, a __main__ block that reports the pass count):

  1. Pure-function tests for the extractors (no disk I/O): the path-citation
     regex and the category-reference / proper-noun extractors, fed
     hand-written strings directly.
  2. End-to-end tests that build a synthetic SVC-PC task package on disk
     (a temp directory with environment/initial_workspace/, solution/files/
     <golden>.xlsx, instruction.md, tests/deliverables/judge.toml) and run
     the full analyze_task against it. These cover the four calibration
     cases the task requires:
       - sub-case A negative: golden cites a real, present file (no finding)
       - sub-case A positive: golden cites a nonexistent file in a cell
         (one finding, kind golden_cites_absent_file, sub_case A)
       - sub-case B positive: instruction names a category with no matching
         environment file (one finding, kind instruction_names_absent_category,
         sub_case B)
       - sub-case B negative: instruction names a category that DOES match a
         real environment filename (no finding)

The sub-case A positive fixture is calibrated against the report-quoted
golden cell content from coverage-cov5-SVC-PC-37-3-0.md (the SwiftShip Reno
MSA path). The two tasks the report cites for this defect (SVC-PC-37-3-0
and SVC-PC-41-4-0) are not among the three real staged tasks on this build
machine, and those three staged tasks carry no solution/ or environment/,
so the SwiftShip calibration here is from the report text quoted into a
synthetic golden xlsx cell, not from a real golden workbook.

Run:
    python deterministic-checks/source_citation_existence_check_test.py
"""

from __future__ import annotations

import shutil
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
if str(HERE) not in sys.path:
    sys.path.insert(0, str(HERE))

from source_citation_existence_check import (  # noqa: E402
    _extract_path_citations,
    _extract_category_references,
    _proper_nouns_in_clause,
    analyze_task,
)

# ---------------------------------------------------------------------------
# Helpers for building synthetic SVC-PC task packages on disk
# ---------------------------------------------------------------------------


def _make_task_dir(
    root: Path,
    task_id: str,
    *,
    golden_cells: dict[str, list[list[str]]] | None = None,
    golden_name: str = "Golden.xlsx",
    env_files: list[str] | None = None,
    instruction: str = "",
) -> Path:
    """Build a minimal SVC-PC task package under root/<task_id>/. Returns the
    task directory path.

    golden_cells: {sheet_name: [[row0cell0, row0cell1], [row1cell0, ...]]}.
    env_files: list of basenames to drop into environment/initial_workspace/.
    """
    task = root / task_id
    (task / "tests" / "deliverables").mkdir(parents=True, exist_ok=True)
    # Minimal judge.toml so detect_task_format() returns "svc_pc" explicitly
    # (it would default to svc_pc anyway, but this makes the format real).
    (task / "tests" / "deliverables" / "judge.toml").write_text(
        '[[criterion]]\nname = "c01"\ndescription = "a test criterion"\n',
        encoding="utf-8",
    )
    env_dir = task / "environment" / "initial_workspace"
    if env_files is not None:
        env_dir.mkdir(parents=True, exist_ok=True)
        for fn in env_files:
            (env_dir / fn).write_text("synthetic environment file", encoding="utf-8")

    if golden_cells is not None:
        import openpyxl

        sol_dir = task / "solution" / "files"
        sol_dir.mkdir(parents=True, exist_ok=True)
        wb = openpyxl.Workbook()
        first = True
        for sheet_name, rows in golden_cells.items():
            ws = wb.active if first else wb.create_sheet()
            ws.title = sheet_name
            for r in rows:
                ws.append(r)
            first = False
        wb.save(sol_dir / golden_name)

    if instruction:
        (task / "instruction.md").write_text(instruction, encoding="utf-8")

    return task


# The report-quoted golden cell content from coverage-cov5-SVC-PC-37-3-0.md
# (RC5/RC6): a path with internal spaces sitting in a single cell, the exact
# shape the prose-focused checks miss.
SWIFISHIP_RENO_PATH = (
    r"Company Shared Drive\Vendors\SwiftShip_Express\Contracts"
    r"\MSA_SwiftShip_2021.pdf"
)


# ---------------------------------------------------------------------------
# Pure-function tests
# ---------------------------------------------------------------------------


def test_whole_cell_path_with_spaces():
    """A path with internal spaces sitting alone in a cell is extracted as
    one 'path' citation (the SwiftShip Reno MSA shape)."""
    res = _extract_path_citations(SWIFISHIP_RENO_PATH)
    assert len(res) == 1, f"expected 1 citation, got {res}"
    citation, bn, kind = res[0]
    assert kind == "path", f"expected path, got {kind}"
    assert bn == "MSA_SwiftShip_2021.pdf", f"wrong basename: {bn}"
    assert "MSA_SwiftShip_2021.pdf" in citation
    print(
        "  test_whole_cell_path_with_spaces: PASS (SwiftShip Reno path -> 1 path citation)"
    )


def test_embedded_path_no_spaces():
    """A path without spaces embedded in prose is extracted as a path."""
    res = _extract_path_citations("see path/to/Revenue.xlsx for details")
    assert len(res) == 1, f"expected 1, got {res}"
    citation, bn, kind = res[0]
    assert kind == "path"
    assert bn == "Revenue.xlsx"
    print("  test_embedded_path_no_spaces: PASS (path/to/Revenue.xlsx extracted)")


def test_bare_filename_in_prose():
    """A bare word.ext token in prose is extracted as a filename citation."""
    res = _extract_path_citations(
        "reconciles to MSA_SwiftShip_2021.pdf per the contract"
    )
    assert len(res) == 1, f"expected 1, got {res}"
    citation, bn, kind = res[0]
    assert kind == "filename"
    assert bn == "MSA_SwiftShip_2021.pdf"
    print("  test_bare_filename_in_prose: PASS (bare MSA_SwiftShip_2021.pdf extracted)")


def test_prose_ending_in_ext_not_whole_chunk():
    """A prose sentence ending in '.pdf' is NOT consumed as a whole-chunk
    path (the sentence break prevents it); only the bare filename is
    extracted."""
    res = _extract_path_citations("The final report was saved as report.pdf.")
    kinds = [k for _c, _b, k in res]
    assert "path" not in kinds, f"should not be a whole-chunk path: {res}"
    assert any(b == "report.pdf" for _c, b, _k in res), f"expected report.pdf: {res}"
    print("  test_prose_ending_in_ext_not_whole_chunk: PASS (prose not consumed whole)")


def test_plain_prose_no_citation():
    """Plain prose with no file-shaped token yields no citation (sub-case B's
    job, not A's)."""
    res = _extract_path_citations("the Apex and SwiftShip rate cards")
    assert res == [], f"expected no citations, got {res}"
    print("  test_plain_prose_no_citation: PASS (no false citation in plain prose)")


def test_category_reference_rate_cards():
    """'the Apex and SwiftShip rate cards' extracts proper nouns apex and
    swiftship against category 'cards' (the cov6 positive)."""
    res = _extract_category_references("Use the Apex and SwiftShip rate cards now.")
    assert len(res) == 1, f"expected 1 ref, got {res}"
    phrase, proper, cat = res[0]
    assert cat == "cards", f"expected cards, got {cat}"
    assert sorted(proper) == ["apex", "swiftship"], f"wrong proper nouns: {proper}"
    print("  test_category_reference_rate_cards: PASS (apex+swiftship extracted)")


def test_category_reference_excludes_months():
    """'the March data book' yields no category reference (March is a month,
    excluded; no proper-noun signal)."""
    res = _extract_category_references("Work from the March data book (v6.2).")
    assert res == [], f"expected no ref (March excluded), got {res}"
    print("  test_category_reference_excludes_months: PASS (March excluded)")


def test_category_reference_no_proper_noun():
    """'Use the report' with no proper noun yields no reference."""
    res = _extract_category_references("Use the report to build the analysis.")
    assert res == [], f"expected no ref, got {res}"
    print("  test_category_reference_no_proper_noun: PASS (no proper noun -> no ref)")


def test_category_reference_determiner_breaks_prose():
    """'Acme Corp filed the report' does NOT extract acme/corp: the determiner
    'the' sits immediately before 'report', breaking the chain before any
    proper noun (the adjacency requirement)."""
    res = _extract_category_references(
        "Acme Corp filed the report. Use the Apex rate cards now."
    )
    # Only the 'cards' reference (apex) should survive; the 'report' one must
    # not, because 'the' breaks the backward walk.
    cats = [c for _p, _pr, c in res]
    assert "report" not in cats, f"'report' should not fire (determiner break): {res}"
    assert "cards" in cats, f"'cards' should fire: {res}"
    print(
        "  test_category_reference_determiner_breaks_prose: PASS (determiner breaks prose)"
    )


def test_proper_nouns_adjacent_chain():
    """Direct test of the adjacency walk for a connective chain."""
    proper = _proper_nouns_in_clause("Use the Apex and SwiftShip rate cards now.", 30)
    assert sorted(proper) == ["apex", "swiftship"], f"got {proper}"
    print("  test_proper_nouns_adjacent_chain: PASS (apex+swiftship via and-chain)")


# ---------------------------------------------------------------------------
# Regression tests: real-corpus false-positive modes (round 32, 27-task
# sample at C:\wt\_shared\svc-task-samples-27-full). Each test uses the exact
# text from a real task that produced a confirmed false positive before the
# fix, not a synthetic example.
# ---------------------------------------------------------------------------


def test_subcase_b_excludes_capitalized_common_words():
    """Real FP from SVC-PC-92-0-0: 'TASK DESIGN DOCUMENT' is a section header,
    not a file reference. design, task, document are common words, not proper
    nouns. No category reference should fire."""
    res = _extract_category_references("TASK DESIGN DOCUMENT")
    assert res == [], f"expected no ref (common capitalized words), got {res}"
    print(
        "  test_subcase_b_excludes_capitalized_common_words: PASS (TASK DESIGN DOCUMENT -> no ref)"
    )


def test_subcase_b_excludes_fy_date_modifier():
    """Real FP from SVC-PC-92-0-0: 'FY2022 Critical Policy' treats FY as a
    fiscal-year date modifier (like March), not a proper noun. critical is
    also a common word. No category reference should fire."""
    res = _extract_category_references(
        "FY2022 Critical Policy and Procedure Lifecycle Controls"
    )
    assert res == [], f"expected no ref (FY + critical excluded), got {res}"
    print(
        "  test_subcase_b_excludes_fy_date_modifier: PASS (FY2022 Critical Policy -> no ref)"
    )


def test_subcase_b_excludes_final_version():
    """Real FP from SVC-PC-13-3-0: 'Use only the named FINAL VERSION source
    files' treats final and version as common words (all-caps emphasis), not
    proper nouns. No category reference should fire for 'files'."""
    res = _extract_category_references("Use only the named FINAL VERSION source files")
    cats = [c for _p, _pr, c in res]
    assert (
        "files" not in cats
    ), f"'files' should not fire (FINAL VERSION are common words): {res}"
    print(
        "  test_subcase_b_excludes_final_version: PASS (FINAL VERSION source files -> no ref)"
    )


def test_subcase_b_excludes_control_gap():
    """Real FP from SVC-PC-14-6-0: 'Control-Gap Register' is a deliverable
    name, not a source reference. control and gap are common words."""
    res = _extract_category_references(
        "Deliverable: a prioritized Control-Gap Register and a Remediation Roadmap."
    )
    cats = [c for _p, _pr, c in res]
    assert (
        "register" not in cats
    ), f"'register' should not fire (control+gap are common words): {res}"
    print(
        "  test_subcase_b_excludes_control_gap: PASS (Control-Gap Register -> no ref)"
    )


def test_subcase_b_excludes_assess_prefer_define():
    """Real FPs: 'Assess and log' (SVC-PC-14-6-0), 'Prefer register'
    (SVC-PC-88-0-0), 'Define document' (SVC-PC-80-7-0). Sentence-start verbs
    capitalized at position 1, not proper nouns."""
    for text, label in [
        (
            "Assess and log design or operating gaps across: financial reporting.",
            "Assess/log",
        ),
        ("Prefer register-versus-GL consistency checks.", "Prefer/register"),
        (
            "Define document and version governance from the organization.",
            "Define/document",
        ),
    ]:
        res = _extract_category_references(text)
        # None of these should produce a finding: the capitalized word is a
        # sentence-start verb, not a proper noun.
        for _p, proper, cat in res:
            assert (
                proper == []
            ), f"{label}: expected no proper nouns, got {proper} for cat={cat}"
    print(
        "  test_subcase_b_excludes_assess_prefer_define: PASS (sentence-start verbs excluded)"
    )


def test_subcase_b_preserves_real_positive_apex_swiftship():
    """The real positive (cov6 SVC-PC-41-4-0: 'the Apex and SwiftShip rate
    cards') must still fire after the COMMON_CAPITALIZED_WORDS fix. apex and
    swiftship are NOT common words."""
    res = _extract_category_references(
        "Use the Apex and SwiftShip rate cards to size the peaks."
    )
    assert len(res) == 1, f"expected 1 ref, got {res}"
    phrase, proper, cat = res[0]
    assert cat == "cards"
    assert sorted(proper) == ["apex", "swiftship"], f"real positive broken: {proper}"
    print(
        "  test_subcase_b_preserves_real_positive_apex_swiftship: PASS (Apex+SwiftShip still fire)"
    )


def test_subcase_b_workday_still_fires():
    """Real case from SVC-PC-88-0-0: 'Workday export' is a real system name
    (Workday Inc.), NOT a common word. It should STILL produce a finding when
    no env file contains 'workday'. This is a structural-limit disclosure, not
    a fix: the export file is named by content, not system name."""
    res = _extract_category_references(
        "The Q4 Workday system export in the data room includes a payroll extract."
    )
    # 'export' is the category word, 'workday' should be extracted as proper noun
    found_workday = False
    for _p, proper, cat in res:
        if cat in ("export", "exports") and "workday" in proper:
            found_workday = True
    assert found_workday, f"workday should still fire (not a common word), got {res}"
    print(
        "  test_subcase_b_workday_still_fires: PASS (Workday still extracted as proper noun)"
    )


def test_subcase_b_common_word_as_modifier_preserves_proper_noun():
    """'the Apex system files' must still extract 'apex' even though 'system'
    is in COMMON_CAPITALIZED_WORDS. 'system' is treated as a modifier, not a
    break, so the walk continues to 'Apex'."""
    res = _extract_category_references(
        "Use the Apex system files to build the analysis."
    )
    assert len(res) >= 1, f"expected at least 1 ref, got {res}"
    found_apex = False
    for _p, proper, cat in res:
        if "apex" in proper:
            found_apex = True
    assert found_apex, f"apex should be reached past 'system' modifier, got {res}"
    print(
        "  test_subcase_b_common_word_as_modifier_preserves_proper_noun: PASS (Apex past system)"
    )


# ---------------------------------------------------------------------------
# End-to-end tests with synthetic on-disk task packages
# ---------------------------------------------------------------------------


def test_subcase_a_negative_cites_present_file():
    """Sub-case A negative: a golden cell cites a file that IS present in the
    environment -> no finding."""
    root = Path(tempfile.mkdtemp(prefix="sce_test_"))
    try:
        _make_task_dir(
            root,
            "SVC-PC-A-NEG-0",
            golden_cells={
                "Evidence": [["source", "see reports\\Q1_Revenue.xlsx for the figure"]]
            },
            env_files=["Q1_Revenue.xlsx", "other.txt"],
            instruction="Build the analysis.\n",
        )
        r = analyze_task(root / "SVC-PC-A-NEG-0")
        a_findings = [f for f in r["findings"] if f["sub_case"] == "A"]
        assert (
            r["status"] == "pass"
        ), f"expected pass, got {r['status']}: {r['findings']}"
        assert a_findings == [], f"expected no sub-case A finding, got {a_findings}"
        print(
            "  test_subcase_a_negative_cites_present_file: PASS (present file -> no finding)"
        )
    finally:
        shutil.rmtree(root, ignore_errors=True)


def test_subcase_a_positive_cites_absent_file():
    """Sub-case A positive: a golden cell cites a file path that is NOT
    present in the environment -> one finding. Calibrated against the
    report-quoted SwiftShip Reno MSA path (cov5)."""
    root = Path(tempfile.mkdtemp(prefix="sce_test_"))
    try:
        _make_task_dir(
            root,
            "SVC-PC-A-POS-0",
            golden_cells={
                "Physical Presence Evidence": [
                    ["claim", "Reno presence per " + SWIFISHIP_RENO_PATH]
                ]
            },
            env_files=["COI_SwiftShipExpress.pdf", "COI_Apex_Logistics.pdf"],
            instruction="Build the analysis.\n",
        )
        r = analyze_task(root / "SVC-PC-A-POS-0")
        a_findings = [f for f in r["findings"] if f["sub_case"] == "A"]
        assert (
            r["status"] == "warn"
        ), f"expected warn, got {r['status']}: {r['findings']}"
        assert len(a_findings) == 1, f"expected 1 sub-case A finding, got {a_findings}"
        f = a_findings[0]
        assert f["kind"] == "golden_cites_absent_file", f"wrong kind: {f['kind']}"
        assert (
            f["basename"] == "MSA_SwiftShip_2021.pdf"
        ), f"wrong basename: {f['basename']}"
        assert f["severity"] == "warn"
        assert (
            "Physical Presence Evidence" in f["location"]
        ), f"wrong location: {f['location']}"
        print(
            "  test_subcase_a_positive_cites_absent_file: PASS (SwiftShip Reno path -> 1 finding)"
        )
    finally:
        shutil.rmtree(root, ignore_errors=True)


def test_subcase_a_skips_golden_own_filename():
    """A golden cell citing the golden's OWN filename (or a sibling golden)
    is not flagged: a golden naming its own output is not a missing source.
    Real calibration for this skip: task_175's golden docx cites
    ipo_valuation_model.xlsx / ipo_pricing_recommendation.docx, its own
    deliverables."""
    root = Path(tempfile.mkdtemp(prefix="sce_test_"))
    try:
        _make_task_dir(
            root,
            "SVC-PC-A-OWN-0",
            golden_cells={
                "Summary": [["see Golden.xlsx and ipo_pricing_recommendation.docx"]]
            },
            golden_name="Golden.xlsx",
            env_files=["only_env_file.txt"],
            instruction="Build the analysis.\n",
        )
        # Add a sibling golden so ipo_pricing_recommendation.docx is also a
        # golden file (its own basename).
        (
            root
            / "SVC-PC-A-OWN-0"
            / "solution"
            / "files"
            / "ipo_pricing_recommendation.docx"
        ).write_text("sibling golden", encoding="utf-8")
        r = analyze_task(root / "SVC-PC-A-OWN-0")
        a_findings = [f for f in r["findings"] if f["sub_case"] == "A"]
        assert (
            a_findings == []
        ), f"golden's own filenames must be skipped, got {a_findings}"
        print(
            "  test_subcase_a_skips_golden_own_filename: PASS (own deliverable names skipped)"
        )
    finally:
        shutil.rmtree(root, ignore_errors=True)


def test_subcase_b_positive_no_matching_env_file():
    """Sub-case B positive: instruction names a category ('rate cards') with
    a proper noun ('Acme') that no environment file contains -> one finding.
    Calibrated against the cov6 SVC-PC-41-4-0 shape ('the Apex and SwiftShip
    rate cards'), using a fresh proper noun so the env match is unambiguous."""
    root = Path(tempfile.mkdtemp(prefix="sce_test_"))
    try:
        _make_task_dir(
            root,
            "SVC-PC-B-POS-0",
            golden_cells={"Out": [["result", "ok"]]},
            env_files=["COI_SwiftShipExpress.pdf", "unrelated_log.csv"],
            instruction=(
                "Build the network design. Use the Acme rate cards to size "
                "the peaks. Return the workbook.\n"
            ),
        )
        r = analyze_task(root / "SVC-PC-B-POS-0")
        b_findings = [f for f in r["findings"] if f["sub_case"] == "B"]
        assert (
            r["status"] == "warn"
        ), f"expected warn, got {r['status']}: {r['findings']}"
        assert len(b_findings) == 1, f"expected 1 sub-case B finding, got {b_findings}"
        f = b_findings[0]
        assert (
            f["kind"] == "instruction_names_absent_category"
        ), f"wrong kind: {f['kind']}"
        assert f["category_word"] == "cards", f"wrong category: {f['category_word']}"
        assert "acme" in f["proper_nouns"], f"missing acme: {f['proper_nouns']}"
        assert f["severity"] == "warn"
        print(
            "  test_subcase_b_positive_no_matching_env_file: PASS (Acme rate cards -> 1 finding)"
        )
    finally:
        shutil.rmtree(root, ignore_errors=True)


def test_subcase_b_negative_category_matches_env_file():
    """Sub-case B negative: instruction names a category ('rate cards') with
    a proper noun ('Apex') that DOES appear in an environment filename
    (Rate_Card_Apex_2021.xlsx) -> no finding (plausible match)."""
    root = Path(tempfile.mkdtemp(prefix="sce_test_"))
    try:
        _make_task_dir(
            root,
            "SVC-PC-B-NEG-0",
            golden_cells={"Out": [["result", "ok"]]},
            env_files=["Rate_Card_Apex_2021.xlsx", "other.txt"],
            instruction=(
                "Build the network design. Use the Apex rate cards to size "
                "the peaks. Return the workbook.\n"
            ),
        )
        r = analyze_task(root / "SVC-PC-B-NEG-0")
        b_findings = [f for f in r["findings"] if f["sub_case"] == "B"]
        assert (
            b_findings == []
        ), f"expected no sub-case B finding (apex matches env), got {b_findings}"
        print(
            "  test_subcase_b_negative_category_matches_env_file: PASS (Apex matches env -> no finding)"
        )
    finally:
        shutil.rmtree(root, ignore_errors=True)


def test_no_environment_dir_is_skip():
    """A task with no environment/ directory reports skip, not error (a real
    coverage gap, not a crash)."""
    root = Path(tempfile.mkdtemp(prefix="sce_test_"))
    try:
        _make_task_dir(
            root,
            "SVC-PC-NOENV-0",
            golden_cells={"Out": [["x", "see missing\\file.pdf"]]},
            env_files=None,
            instruction="Build.\n",
        )
        r = analyze_task(root / "SVC-PC-NOENV-0")
        assert r["status"] == "skip", f"expected skip, got {r['status']}"
        assert r["n_findings"] == 0
        print("  test_no_environment_dir_is_skip: PASS (no env -> skip, not error)")
    finally:
        shutil.rmtree(root, ignore_errors=True)


def test_subcase_a_excludes_wildcard_hash_basename():
    """Real FP from SVC-PC-14-1-0: golden cell 'Q#_Board_Minutes.docx' uses #
    as a wildcard (Q1/Q2/Q3/Q4_Board_Minutes.docx). No env file has # in its
    name; the citation is a template, not a missing source."""
    root = Path(tempfile.mkdtemp(prefix="sce_test_"))
    try:
        _make_task_dir(
            root,
            "SVC-PC-A-HASH-0",
            golden_cells={
                "Master_PBC_Tracker": [
                    [
                        "evidence",
                        "Company Shared Drive/Board/Meetings/2021-Q1..Q4/Q#_Board_Minutes.docx",
                    ]
                ]
            },
            env_files=["Q1_Board_Minutes.docx", "Q2_Board_Minutes.docx"],
            instruction="Build the analysis.\n",
        )
        r = analyze_task(root / "SVC-PC-A-HASH-0")
        a_findings = [f for f in r["findings"] if f["sub_case"] == "A"]
        hash_findings = [f for f in a_findings if "#" in f["basename"]]
        assert (
            hash_findings == []
        ), f"Q# wildcard should be filtered, got {hash_findings}"
        print(
            "  test_subcase_a_excludes_wildcard_hash_basename: PASS (Q#_Board_Minutes.docx filtered)"
        )
    finally:
        shutil.rmtree(root, ignore_errors=True)


def test_subcase_a_excludes_parenthetical_basename():
    """Real FP from SVC-PC-14-1-0 G6: golden cell contains a parenthetical
    list of filenames mangled into one basename by the whole-cell path pass.
    The basename starts with '(' or whitespace -> extraction artifact."""
    root = Path(tempfile.mkdtemp(prefix="sce_test_"))
    try:
        _make_task_dir(
            root,
            "SVC-PC-A-PAREN-0",
            golden_cells={
                "Master_PBC_Tracker": [
                    [
                        "legal",
                        "Company Shared Drive/Legal/ (D&O_Insurance_2021, General_Liability_Policy_2021, Workers_Comp_2021, Property_Insurance_2021, Cyber_Liability_2021).pdf",
                    ]
                ]
            },
            env_files=["D&O_Insurance_2021.pdf"],
            instruction="Build the analysis.\n",
        )
        r = analyze_task(root / "SVC-PC-A-PAREN-0")
        a_findings = [f for f in r["findings"] if f["sub_case"] == "A"]
        paren_findings = [
            f
            for f in a_findings
            if "(" in f["basename"] or f["basename"] != f["basename"].strip()
        ]
        assert (
            paren_findings == []
        ), f"parenthetical basename should be filtered, got {paren_findings}"
        print(
            "  test_subcase_a_excludes_parenthetical_basename: PASS (parenthetical list filtered)"
        )
    finally:
        shutil.rmtree(root, ignore_errors=True)


def test_subcase_a_excludes_bare_year_month():
    """Real FP from SVC-PC-14-1-0 G11: golden cell 'Bank_Statement_2021-01.pdf
    .. 2021-12.pdf (all 12 present)'. The descriptive citation is matched
    correctly; the bare '2021-12.pdf' range-end shorthand is filtered."""
    root = Path(tempfile.mkdtemp(prefix="sce_test_"))
    try:
        _make_task_dir(
            root,
            "SVC-PC-A-YM-0",
            golden_cells={
                "Master_PBC_Tracker": [
                    [
                        "bank",
                        "Company Shared Drive/Accounting/2021/Bank_Statements/Bank_Statement_2021-01.pdf .. 2021-12.pdf (all 12 present)",
                    ]
                ]
            },
            env_files=["Bank_Statement_2021-01.pdf", "Bank_Statement_2021-12.pdf"],
            instruction="Build the analysis.\n",
        )
        r = analyze_task(root / "SVC-PC-A-YM-0")
        a_findings = [f for f in r["findings"] if f["sub_case"] == "A"]
        ym_findings = [
            f for f in a_findings if f["basename"] in ("2021-12.pdf", "2021-12.xlsx")
        ]
        assert (
            ym_findings == []
        ), f"bare year-month should be filtered, got {ym_findings}"
        print(
            "  test_subcase_a_excludes_bare_year_month: PASS (2021-12.pdf range shorthand filtered)"
        )
    finally:
        shutil.rmtree(root, ignore_errors=True)


def test_subcase_a_real_positive_cites_present_file():
    """Real positive from SVC-PC-14-1-0: the golden PBC tracker cites 100+
    real environment files by exact basename (Bank_Statement_2021-01.pdf,
    Q1_Bank_Reconciliation.xlsx, etc.). Sub-case A correctly stays silent on
    all matched citations. This is the real-positive calibration the round-31
    review said was never tested."""
    root = Path(tempfile.mkdtemp(prefix="sce_test_"))
    try:
        _make_task_dir(
            root,
            "SVC-PC-A-REAL-0",
            golden_cells={
                "Master_PBC_Tracker": [
                    [
                        "desc",
                        "Company Shared Drive/Accounting/2021/Bank_Statements/Bank_Statement_2021-01.pdf",
                    ],
                    [
                        "desc2",
                        "Company Shared Drive/Accounting/2021/Reconciliations/Q1_Bank_Reconciliation.xlsx",
                    ],
                ]
            },
            env_files=[
                "Bank_Statement_2021-01.pdf",
                "Q1_Bank_Reconciliation.xlsx",
                "Chart_of_Accounts_.xlsx",
            ],
            instruction="Build the analysis.\n",
        )
        r = analyze_task(root / "SVC-PC-A-REAL-0")
        a_findings = [f for f in r["findings"] if f["sub_case"] == "A"]
        assert (
            a_findings == []
        ), f"real present files should produce no A finding, got {a_findings}"
        print(
            "  test_subcase_a_real_positive_cites_present_file: PASS (Bank_Statement + Q1_Reconciliation matched)"
        )
    finally:
        shutil.rmtree(root, ignore_errors=True)


def test_subcase_a_excludes_underscore_suffix():
    """Real FP from SVC-PC-14-1-0: '_v2.docx' and '_DOCX_SOURCE.docx' are
    suffix fragments split from longer filenames (Board_Composition_v2.docx,
    FY2021_CS_and_Returns_Report_DOCX_SOURCE.docx). Real filenames never
    start with '_'."""
    root = Path(tempfile.mkdtemp(prefix="sce_test_"))
    try:
        _make_task_dir(
            root,
            "SVC-PC-A-UND-0",
            golden_cells={
                "Evidence_Index": [
                    [
                        "doc",
                        "Company Shared Drive/Legal/Board_Composition_v2.docx and related _v2.docx",
                    ],
                ]
            },
            env_files=["Board_Composition_v2.docx"],
            instruction="Build the analysis.\n",
        )
        r = analyze_task(root / "SVC-PC-A-UND-0")
        a_findings = [f for f in r["findings"] if f["sub_case"] == "A"]
        underscore = [f for f in a_findings if f["basename"].startswith("_")]
        assert (
            underscore == []
        ), f"underscore-suffix basenames should be filtered, got {underscore}"
        print("  test_subcase_a_excludes_underscore_suffix: PASS (_v2.docx filtered)")
    finally:
        shutil.rmtree(root, ignore_errors=True)


def test_subcase_a_excludes_short_digit_stem():
    """Real FP from SVC-PC-14-1-0: '12.pdf' and '03.pdf' are date/number
    fragments, not real filenames. Real audit files always have descriptive
    prefixes."""
    root = Path(tempfile.mkdtemp(prefix="sce_test_"))
    try:
        _make_task_dir(
            root,
            "SVC-PC-A-DIG-0",
            golden_cells={
                "Evidence_Index": [
                    [
                        "inv",
                        "Company Shared Drive/Vendors/Invoice_Apex_2021-12.pdf and 12.pdf",
                    ],
                    [
                        "po",
                        "Company Shared Drive/Vendors/NovaTex_PO_2021_Q3_003.pdf and 03.pdf",
                    ],
                ]
            },
            env_files=["Invoice_Apex_2021-12.pdf"],
            instruction="Build the analysis.\n",
        )
        r = analyze_task(root / "SVC-PC-A-DIG-0")
        a_findings = [f for f in r["findings"] if f["sub_case"] == "A"]
        short_digits = [
            f
            for f in a_findings
            if "." in f["basename"]
            and f["basename"].rsplit(".", 1)[0].isdigit()
            and len(f["basename"].rsplit(".", 1)[0]) < 4
        ]
        assert (
            short_digits == []
        ), f"short digit-only stems should be filtered, got {short_digits}"
        print(
            "  test_subcase_a_excludes_short_digit_stem: PASS (12.pdf, 03.pdf filtered)"
        )
    finally:
        shutil.rmtree(root, ignore_errors=True)


def test_subcase_a_excludes_semicolon_multifile():
    """Real FP from SVC-PC-14-1-0: cell value 'Q4_Processing_Statement.pdf;
    Annual_Processing_Summary.pdf' is two filenames joined by ';', not one
    citation."""
    root = Path(tempfile.mkdtemp(prefix="sce_test_"))
    try:
        _make_task_dir(
            root,
            "SVC-PC-A-SEM-0",
            golden_cells={
                "Evidence_Index": [
                    [
                        "proc",
                        "Company Shared Drive/Vendors/Q4_Processing_Statement.pdf; Annual_Processing_Summary.pdf",
                    ],
                ]
            },
            env_files=["Q4_Processing_Statement.pdf"],
            instruction="Build the analysis.\n",
        )
        r = analyze_task(root / "SVC-PC-A-SEM-0")
        a_findings = [f for f in r["findings"] if f["sub_case"] == "A"]
        semis = [f for f in a_findings if ";" in f["basename"]]
        assert (
            semis == []
        ), f"semicolon-joined basenames should be filtered, got {semis}"
        print(
            "  test_subcase_a_excludes_semicolon_multifile: PASS (multi-file ';' filtered)"
        )
    finally:
        shutil.rmtree(root, ignore_errors=True)


def test_subcase_b_excludes_mdf_okr_swot_sor():
    """Real FPs from the 202-task corpus: MDF (Market Development Funds,
    SVC-PC-35-0-0), OKR (Objectives & Key Results, SVC-PC-53-6-0), SWOT
    (analysis framework, SVC-PC-8-7-0), SOR (System of Record, SVC-PC-80-6-0)
    are business acronyms, not file-identifying proper nouns."""
    for text, pn in [
        ("4. MDF governance policy", "mdf"),
        ("a complete FY2022 customer support OKR plan", "okr"),
        ("company SWOT, and risk register", "swot"),
        ("so they can slot into the SOR policy", "sor"),
    ]:
        res = _extract_category_references(text)
        for _p, proper, _c in res:
            assert pn not in proper, f"{pn} should not be a proper noun: {proper}"
    print(
        "  test_subcase_b_excludes_mdf_okr_swot_sor: PASS (business acronyms excluded)"
    )


def test_subcase_b_excludes_enterprise_exposure_insurable():
    """Real FP from SVC-PC-25-1-0: 'FY2021 Enterprise Insurable-Exposure
    Register' is a deliverable name, not a source reference. enterprise,
    exposure, insurable are title-case adjectives/nouns, not proper nouns."""
    res = _extract_category_references(
        "Build the FY2021 Enterprise Insurable-Exposure Register."
    )
    for _p, proper, cat in res:
        assert "enterprise" not in proper, f"enterprise should be excluded: {proper}"
        assert "exposure" not in proper, f"exposure should be excluded: {proper}"
        assert "insurable" not in proper, f"insurable should be excluded: {proper}"
    print(
        "  test_subcase_b_excludes_enterprise_exposure_insurable: PASS (title-case words excluded)"
    )


def test_subcase_a_excludes_companion_golden_reference():
    """Real FP from SVC-PC-7-4-0 (round 42): a two-golden (companion-
    deliverable) task where each golden cross-references the OTHER by an
    internal/pre-rename name ("L2-04_Task_7_Golden.xlsx") instead of its
    shipped basename. skip_basenames only knows shipped names; the
    dedicated 'golden' substring + 2-golden-files gate catches this."""
    root = Path(tempfile.mkdtemp(prefix="sce_test_"))
    try:
        _make_task_dir(
            root,
            "SVC-PC-A-GOLDEN-0",
            golden_cells={
                "README": [
                    ["title", "Ops Model GOLDEN"],
                    ["note", "Companion summary workbook is L2-04_Task_7_Golden.xlsx."],
                ]
            },
            golden_name="Cirrus_Ops_Model.xlsx",
            env_files=["_no_input_files_required.txt"],
            instruction="Build the analysis.\n",
        )
        # Add the second, companion golden directly (the helper only builds
        # one golden file per call).
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Recommendation"
        ws.append(["note", "Companion ops model: Cirrus_Sleep_Ops_Model_GOLDEN.xlsx."])
        wb.save(
            root
            / "SVC-PC-A-GOLDEN-0"
            / "solution"
            / "files"
            / "Cirrus_Vendor_Plan.xlsx"
        )

        r = analyze_task(root / "SVC-PC-A-GOLDEN-0")
        a_findings = [f for f in r["findings"] if f["sub_case"] == "A"]
        golden_findings = [f for f in a_findings if "golden" in f["basename"].lower()]
        assert (
            golden_findings == []
        ), f"companion-golden cross-reference should be filtered, got {golden_findings}"
        print(
            "  test_subcase_a_excludes_companion_golden_reference: PASS (companion-golden internal name filtered)"
        )
    finally:
        shutil.rmtree(root, ignore_errors=True)


def test_subcase_a_single_golden_task_golden_substring_still_flags():
    """Guard against over-suppression (round 42): a SINGLE-golden task
    whose golden cites a basename containing 'golden' is NOT covered by
    the companion-deliverable filter (gated on 2+ golden files) and must
    still be flagged. Real corpus check: SVC-PC-3-4-0, SVC-PC-7-2-0,
    SVC-PC-94-3-0 each have exactly 1 golden file and real
    'golden'-substring findings that the round-42 fix must NOT silently
    suppress (9 findings across those 3 tasks, confirmed untouched by a
    real full-600-task before/after diff)."""
    root = Path(tempfile.mkdtemp(prefix="sce_test_"))
    try:
        _make_task_dir(
            root,
            "SVC-PC-A-SINGLEGOLD-0",
            golden_cells={
                "Source_Register": [
                    ["src", "See Legacy_Golden_Copy.xlsx for archived figures."]
                ]
            },
            env_files=["Chart_of_Accounts.xlsx"],
            instruction="Build the analysis.\n",
        )
        r = analyze_task(root / "SVC-PC-A-SINGLEGOLD-0")
        a_findings = [f for f in r["findings"] if f["sub_case"] == "A"]
        golden_findings = [f for f in a_findings if "golden" in f["basename"].lower()]
        assert (
            len(golden_findings) == 1
        ), f"single-golden task's 'golden'-substring citation should still be flagged, got {golden_findings}"
        print(
            "  test_subcase_a_single_golden_task_golden_substring_still_flags: "
            "PASS (1-golden task not over-suppressed)"
        )
    finally:
        shutil.rmtree(root, ignore_errors=True)


def test_subcase_a_excludes_version_fragment_tail():
    """Real FP from SVC-PC-80-7-0 and SVC-PC-80-5-0 (round 42): a golden's
    own naming-convention EXAMPLE ("..._v2.0_FINAL.docx") has its tail
    matched as a bare filename ("0_FINAL.docx") because the regex sees a
    word boundary right after the "." in "v2.0". Filtered by rejecting
    any bare-filename match immediately preceded by "<digit>."."""
    root = Path(tempfile.mkdtemp(prefix="sce_test_"))
    try:
        _make_task_dir(
            root,
            "SVC-PC-A-VERFRAG-0",
            golden_cells={
                "Spec": [
                    ["rule", "Example: Cirrus_Finance_Estimate_FY2021_v2.0_FINAL.docx"]
                ]
            },
            env_files=["Chart_of_Accounts.xlsx"],
            instruction="Build the analysis.\n",
        )
        r = analyze_task(root / "SVC-PC-A-VERFRAG-0")
        a_findings = [f for f in r["findings"] if f["sub_case"] == "A"]
        frag_findings = [f for f in a_findings if f["basename"] == "0_FINAL.docx"]
        assert (
            frag_findings == []
        ), f"version-fragment tail should be filtered, got {frag_findings}"
        print(
            "  test_subcase_a_excludes_version_fragment_tail: PASS (0_FINAL.docx tail-split filtered)"
        )
    finally:
        shutil.rmtree(root, ignore_errors=True)


def test_subcase_a_quarter_range_satisfied_suppresses():
    """Real FP from SVC-PC-13-3-0 / SVC-PC-42-1-0 (round 42): golden cites
    'Q1-Q4_Channel_Attribution.xlsx' as range shorthand for 4 separate
    quarterly files, all of which are shipped. Suppressed only because ALL
    4 are confirmed present in the environment."""
    root = Path(tempfile.mkdtemp(prefix="sce_test_"))
    try:
        _make_task_dir(
            root,
            "SVC-PC-A-QRANGE-0",
            golden_cells={
                "Source_Register": [["src", "Q1-Q4_Channel_Attribution.xlsx"]]
            },
            env_files=[
                "Q1_Channel_Attribution.xlsx",
                "Q2_Channel_Attribution.xlsx",
                "Q3_Channel_Attribution.xlsx",
                "Q4_Channel_Attribution.xlsx",
            ],
            instruction="Build the analysis.\n",
        )
        r = analyze_task(root / "SVC-PC-A-QRANGE-0")
        a_findings = [f for f in r["findings"] if f["sub_case"] == "A"]
        range_findings = [
            f
            for f in a_findings
            if f["basename"].lower() == "q1-q4_channel_attribution.xlsx"
        ]
        assert (
            range_findings == []
        ), f"fully-satisfied Q1-Q4 range should be suppressed, got {range_findings}"
        print(
            "  test_subcase_a_quarter_range_satisfied_suppresses: "
            "PASS (Q1-Q4 range with all 4 quarters present suppressed)"
        )
    finally:
        shutil.rmtree(root, ignore_errors=True)


def test_subcase_a_quarter_range_partial_still_flags():
    """Conservative-by-construction check (round 42): a Q1-Q4 range
    citation with only 3 of 4 quarters shipped is a genuine gap, not
    shorthand noise, and must stay flagged."""
    root = Path(tempfile.mkdtemp(prefix="sce_test_"))
    try:
        _make_task_dir(
            root,
            "SVC-PC-A-QPARTIAL-0",
            golden_cells={
                "Source_Register": [["src", "Q1-Q4_Channel_Attribution.xlsx"]]
            },
            env_files=[
                "Q1_Channel_Attribution.xlsx",
                "Q2_Channel_Attribution.xlsx",
                "Q3_Channel_Attribution.xlsx",
                # Q4 deliberately missing: the range is only partially shipped.
            ],
            instruction="Build the analysis.\n",
        )
        r = analyze_task(root / "SVC-PC-A-QPARTIAL-0")
        a_findings = [f for f in r["findings"] if f["sub_case"] == "A"]
        range_findings = [
            f
            for f in a_findings
            if f["basename"].lower() == "q1-q4_channel_attribution.xlsx"
        ]
        assert (
            len(range_findings) == 1
        ), f"partial Q1-Q4 range (Q4 missing) should stay flagged, got {range_findings}"
        print(
            "  test_subcase_a_quarter_range_partial_still_flags: PASS (partial range, Q4 missing, still flagged)"
        )
    finally:
        shutil.rmtree(root, ignore_errors=True)


if __name__ == "__main__":
    tests = [
        test_whole_cell_path_with_spaces,
        test_embedded_path_no_spaces,
        test_bare_filename_in_prose,
        test_prose_ending_in_ext_not_whole_chunk,
        test_plain_prose_no_citation,
        test_category_reference_rate_cards,
        test_category_reference_excludes_months,
        test_category_reference_no_proper_noun,
        test_category_reference_determiner_breaks_prose,
        test_proper_nouns_adjacent_chain,
        test_subcase_b_excludes_capitalized_common_words,
        test_subcase_b_excludes_fy_date_modifier,
        test_subcase_b_excludes_final_version,
        test_subcase_b_excludes_control_gap,
        test_subcase_b_excludes_assess_prefer_define,
        test_subcase_b_preserves_real_positive_apex_swiftship,
        test_subcase_b_workday_still_fires,
        test_subcase_b_common_word_as_modifier_preserves_proper_noun,
        test_subcase_a_negative_cites_present_file,
        test_subcase_a_positive_cites_absent_file,
        test_subcase_a_skips_golden_own_filename,
        test_subcase_b_positive_no_matching_env_file,
        test_subcase_b_negative_category_matches_env_file,
        test_subcase_a_excludes_wildcard_hash_basename,
        test_subcase_a_excludes_parenthetical_basename,
        test_subcase_a_excludes_bare_year_month,
        test_subcase_a_real_positive_cites_present_file,
        test_subcase_a_excludes_underscore_suffix,
        test_subcase_a_excludes_short_digit_stem,
        test_subcase_a_excludes_semicolon_multifile,
        test_subcase_b_excludes_mdf_okr_swot_sor,
        test_subcase_b_excludes_enterprise_exposure_insurable,
        test_no_environment_dir_is_skip,
        test_subcase_a_excludes_companion_golden_reference,
        test_subcase_a_single_golden_task_golden_substring_still_flags,
        test_subcase_a_excludes_version_fragment_tail,
        test_subcase_a_quarter_range_satisfied_suppresses,
        test_subcase_a_quarter_range_partial_still_flags,
    ]
    passed = 0
    failed = 0
    for t in tests:
        try:
            t()
            passed += 1
        except Exception as exc:
            failed += 1
            print(f"  {t.__name__}: FAIL - {type(exc).__name__}: {exc}")
    print(f"\n{passed}/{passed + failed} tests passed")
    sys.exit(1 if failed else 0)
