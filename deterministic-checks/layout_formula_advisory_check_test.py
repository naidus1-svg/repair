#!/usr/bin/env python3
"""Tests for layout_formula_advisory_check.py (G1 / H1 / H2 / I1) and the I2 fix text on reward_signal_gate G6.

Small synthetic fixtures, one positive and one negative per rule. Every rule is advisory: a positive
must come back ``warn``, never ``fail``.
"""

from __future__ import annotations

import json
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
if str(HERE) not in sys.path:
    sys.path.insert(0, str(HERE))

from layout_formula_advisory_check import check_task  # noqa: E402
import reward_signal_gate as RSG  # noqa: E402

FREE_NAMING = "Structure (any sheet/tab naming or count is fine as long as every field exists). Build the plan.\n"


def _det(
    name: str, command: str, args: dict, path: str, comparison: str, expected
) -> dict:
    return {
        "name": name,
        "source": {
            "type": "file",
            "file": {"type": "xlsx", "command": command, "arguments": args},
        },
        "assertion": {
            "type": "deterministic",
            "expected": expected,
            "deterministic": {"path": path, "comparison": comparison},
        },
        "metadata": {"weight": 0.05},
    }


def _native(
    root: Path,
    instruction: str,
    verifiers: list[dict],
    golden_formulas: int | None = None,
) -> Path:
    pkg = root / "task_X" / "rev_1__1" / "task_files"
    (pkg / "tests").mkdir(parents=True)
    (pkg / "instruction.md").write_text(instruction, encoding="utf-8")
    (pkg / "tests" / "verifier.json").write_text(
        json.dumps({"verifiers": verifiers, "config": {}}), encoding="utf-8"
    )
    if golden_formulas is not None:
        from openpyxl import Workbook

        gold = pkg / "solution" / "golden_outputs"
        gold.mkdir(parents=True)
        wb = Workbook()
        ws = wb.active
        for i in range(1, golden_formulas + 1):
            ws.cell(row=i, column=1, value=i)
            ws.cell(row=i, column=2, value=f"=A{i}*2")
        wb.save(gold / "plan.xlsx")
    return root / "task_X"


def _rewardkit(root: Path, criterion: str, test_sh: str, files: list[str]) -> Path:
    pkg = root / "SVC-1" / "rev_1__1" / "task_files"
    grp = pkg / "tests" / "deliverables_01"
    grp.mkdir(parents=True)
    (pkg / "environment" / "data").mkdir(parents=True)
    (pkg / "environment" / "data" / "input.xlsx").write_bytes(b"x")
    (pkg / "instruction.md").write_text("Build a model workbook.\n", encoding="utf-8")
    (pkg / "tests" / "test.sh").write_text(test_sh, encoding="utf-8")
    files_toml = ", ".join(f'"{f}"' for f in files)
    (grp / "judge.toml").write_text(
        f'[judge]\njudge = "anthropic/claude-opus-5"\nfiles = [{files_toml}]\n\n'
        f'[[criterion]]\nname = "c1"\ndescription = "{criterion}"\nweight = 1.0\n',
        encoding="utf-8",
    )
    return root / "SVC-1"


# ----------------------------------------------------------------------------- G1


def test_g1_sheet_pin_under_free_naming_warns() -> None:
    with tempfile.TemporaryDirectory() as td:
        v = _det(
            "d4_named",
            "formula_authenticity_check",
            {"path": "output/plan.xlsx", "sheet": "3. Product", "cell_range": "G4:G7"},
            "$.n_hardcoded",
            "equals",
            0,
        )
        r = check_task(_native(Path(td), FREE_NAMING, [v]))
        g1 = r["checks"]["G1_layout_pin"]
        assert g1["status"] == "warn", g1
        assert (
            "d4_named" in g1["evidence"][0]
            and "tests/verifier.json" in g1["evidence"][0]
            and "Fix:" in g1["evidence"][0]
        )
        assert r["status"] == "warn"


def test_g1_sheet_named_in_instruction_passes() -> None:
    with tempfile.TemporaryDirectory() as td:
        v = _det(
            "d4_named",
            "formula_authenticity_check",
            {"path": "output/plan.xlsx", "sheet": "3. Product", "cell_range": "G4:G7"},
            "$.n_hardcoded",
            "equals",
            0,
        )
        instr = FREE_NAMING + "The tab '3. Product' must hold formulas in G4:G7.\n"
        assert (
            check_task(_native(Path(td), instr, [v]))["checks"]["G1_layout_pin"][
                "status"
            ]
            == "pass"
        )


def test_g1_no_free_naming_passes() -> None:
    with tempfile.TemporaryDirectory() as td:
        v = _det(
            "d4_named",
            "formula_authenticity_check",
            {"path": "output/plan.xlsx", "sheet": "Model", "cell_range": "A1"},
            "$.n_hardcoded",
            "equals",
            0,
        )
        assert (
            check_task(_native(Path(td), "Use a tab called Model.\n", [v]))["checks"][
                "G1_layout_pin"
            ]["status"]
            == "pass"
        )


# ----------------------------------------------------------------------------- H1


def test_h1_floor_equal_to_golden_warns() -> None:
    with tempfile.TemporaryDirectory() as td:
        v = _det(
            "d2_counts",
            "formula_summary",
            {"path": "output/plan.xlsx"},
            "$.n_formula_with_refs",
            "greater_than_equal",
            6,
        )
        h1 = check_task(
            _native(Path(td), "Counts must be live formulas.\n", [v], golden_formulas=6)
        )["checks"]["H1_formula_floor"]
        assert h1["status"] == "warn", h1
        assert (
            "d2_counts" in h1["evidence"][0]
            and "has 6" in h1["evidence"][0]
            and "Fix:" in h1["evidence"][0]
        )


def test_h1_floor_above_golden_warns() -> None:
    with tempfile.TemporaryDirectory() as td:
        v = _det(
            "d2_counts",
            "formula_summary",
            {"path": "output/plan.xlsx"},
            "$.n_formula_with_refs",
            "greater_than_equal",
            9,
        )
        h1 = check_task(
            _native(Path(td), "Counts must be live formulas.\n", [v], golden_formulas=6)
        )["checks"]["H1_formula_floor"]
        assert h1["status"] == "warn" and "golden itself fails" in h1["evidence"][0]


def test_h1_floor_with_headroom_passes() -> None:
    with tempfile.TemporaryDirectory() as td:
        v = _det(
            "d2_counts",
            "formula_summary",
            {"path": "output/plan.xlsx"},
            "$.n_formula_with_refs",
            "greater_than_equal",
            3,
        )
        h1 = check_task(
            _native(Path(td), "Counts must be live formulas.\n", [v], golden_formulas=6)
        )["checks"]["H1_formula_floor"]
        assert h1["status"] == "pass", h1


# ----------------------------------------------------------------------------- H2


def test_h2_formula_gate_without_formula_ask_warns() -> None:
    with tempfile.TemporaryDirectory() as td:
        v = _det(
            "d2_share",
            "formula_authenticity_scan",
            {"path": "output/tracker.xlsx", "sheet_name_contains": []},
            "$.formula_share_all",
            "greater_than_equal",
            0.15,
        )
        instr = "Write an action tracker with columns for owner, priority, timing and account notes.\n"
        h2 = check_task(_native(Path(td), instr, [v]))["checks"]["H2_formula_unasked"]
        assert h2["status"] == "warn", h2
        assert "d2_share" in h2["evidence"][0] and "Fix:" in h2["evidence"][0]


def test_h2_formula_asked_passes() -> None:
    with tempfile.TemporaryDirectory() as td:
        v = _det(
            "d2_share",
            "formula_authenticity_scan",
            {"path": "output/tracker.xlsx", "sheet_name_contains": []},
            "$.formula_share_all",
            "greater_than_equal",
            0.15,
        )
        instr = "Totals must be live formulas referencing the input cells.\n"
        assert (
            check_task(_native(Path(td), instr, [v]))["checks"]["H2_formula_unasked"][
                "status"
            ]
            == "pass"
        )


def test_h2_word_boundaries_do_not_match_account_or_deliverable() -> None:
    with tempfile.TemporaryDirectory() as td:
        v = _det(
            "d2_share",
            "formula_authenticity_scan",
            {"path": "output/t.xlsx", "sheet_name_contains": []},
            "$.formula_share_all",
            "greater_than_equal",
            0.15,
        )
        instr = "Deliverables: an account tracker for delivery owners.\n"
        assert (
            check_task(_native(Path(td), instr, [v]))["checks"]["H2_formula_unasked"][
                "status"
            ]
            == "warn"
        )


# ----------------------------------------------------------------------------- I1


def test_i1_formula_criterion_without_recalc_warns() -> None:
    with tempfile.TemporaryDirectory() as td:
        r = check_task(
            _rewardkit(
                Path(td),
                "The totals row is a live formula summing the lines.",
                "#!/bin/bash\nrewardkit run\n",
                ["/app/output/model.xlsx"],
            )
        )
        i1 = r["checks"]["I1_rk_no_recalc"]
        assert i1["status"] == "warn", i1
        assert (
            "judge.toml" in i1["evidence"][0]
            and "c1" in i1["evidence"][0]
            and "recalc" in i1["evidence"][0]
        )


def test_i1_recalc_step_present_passes() -> None:
    with tempfile.TemporaryDirectory() as td:
        sh = "#!/bin/bash\nsoffice --headless --convert-to xlsx /app/output/model.xlsx\nrewardkit run\n"
        r = check_task(
            _rewardkit(
                Path(td),
                "The totals row is a live formula.",
                sh,
                ["/app/output/model.xlsx"],
            )
        )
        assert r["checks"]["I1_rk_no_recalc"]["status"] == "pass"


def test_i1_input_workbook_only_passes() -> None:
    with tempfile.TemporaryDirectory() as td:
        r = check_task(
            _rewardkit(
                Path(td),
                "Uses the formula from the input.",
                "#!/bin/bash\n",
                ["/app/data/input.xlsx", "/app/output/memo.docx"],
            )
        )
        assert r["checks"]["I1_rk_no_recalc"]["status"] == "pass"


def test_i1_skips_native_task_and_rules_never_fail() -> None:
    with tempfile.TemporaryDirectory() as td:
        v = _det(
            "d2_share",
            "formula_authenticity_scan",
            {"path": "output/t.xlsx", "sheet_name_contains": []},
            "$.formula_share_all",
            "greater_than_equal",
            0.15,
        )
        r = check_task(_native(Path(td), FREE_NAMING, [v], golden_formulas=1))
        assert r["checks"]["I1_rk_no_recalc"]["status"] == "skip"
        assert all(c["status"] != "fail" for c in r["checks"].values())


# ----------------------------------------------------------------------------- I2 (reward_signal_gate G6 fix text)


def _lib(root: Path, body: str) -> Path:
    pkg = root / "task_files"
    src = pkg / "tests" / "verifier_lib" / "sources"
    src.mkdir(parents=True)
    (src / "xlsx.py").write_text(body, encoding="utf-8")
    return pkg


def test_i2_drifted_verifier_lib_names_fix() -> None:
    with tempfile.TemporaryDirectory() as td:
        g6 = RSG.g6_uncached(
            _lib(Path(td), "def extract_xlsx(path, limit):\n    pass\n"),
            [],
            None,
            "task_X",
        )
        assert g6["status"] == "fail"
        assert any(
            "re-vendor verifier_lib" in e and "tests/verifier_lib/sources/xlsx.py" in e
            for e in g6["evidence"]
        ), g6


def test_i2_current_verifier_lib_passes() -> None:
    with tempfile.TemporaryDirectory() as td:
        g6 = RSG.g6_uncached(
            _lib(
                Path(td), "def _recalculated_copy_if_uncached(path):\n    return path\n"
            ),
            [],
            None,
            "task_X",
        )
        assert g6["status"] == "pass" and not any(
            "re-vendor" in e for e in g6["evidence"]
        )


if __name__ == "__main__":
    tests = [
        v for k, v in dict(globals()).items() if k.startswith("test_") and callable(v)
    ]
    for test in tests:
        test()
        print(f"{test.__name__}: PASS")
