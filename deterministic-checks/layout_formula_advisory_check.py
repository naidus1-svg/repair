#!/usr/bin/env python3
"""Layout / formula-gate advisory check: four mechanical rules from the QC batch 1-6 failure analysis.

Source: failure-pattern analysis of 60 failed tasks from QC batches 1-6 (2026-09-28, 169 locally staged
tasks). Each rule caught a recorded failure; each has only 1 to 3 true positives, so EVERY rule here is
advisory (status ``warn``), never ``fail``. Do not promote one to blocking without measured precision.

  G1 layout_pin        EKW-native deterministic verifier passes a ``sheet`` or ``cell_range`` argument
                       while the instruction lets the solver name/count tabs freely ("any sheet/tab naming
                       or count is fine"). A compliant layout fails the check. Caught GTM-178 d4/d5.
  H1 formula_floor     deterministic floor on a formula count/share (``$.n_formula_with_refs``,
                       ``$.n_formula``, ``$.formula_share_all`` ...) at or above the golden workbook's own
                       value: zero headroom, so any compliant layout with one fewer formula fails.
                       Caught HLT-189 d2 (floor 6, golden has exactly 6).
  H2 formula_unasked   formula-count/share verifier on a task whose instruction never asks for formulas
                       (or live/linked/calculated cells). Caught GTM-187 d2 (15% share on a text tracker).
  I1 rk_no_recalc      rewardkit task (judge.toml) whose judge group reads a deliverable .xlsx and whose
                       criteria mention formulas/live/computed values, with no soffice/libreoffice/recalc
                       step in tests/test.sh or environment/Dockerfile: the judge sees blank cached values
                       for an agent's uncached formulas. Caught num-562; present in ~50 of 130 rewardkit
                       tasks, so advisory only.

(I2, a vendored verifier_lib xlsx.py without the uncached-formula recalc, is reward_signal_gate.py G6.)

Output per task: ``{"status": "pass"|"warn"|"skip", "task", "package", "checks": {rule: {status, evidence}}}``.
Every evidence line names the file, the verifier/criterion id, and the fix.

Run: layout_formula_advisory_check.py --root <task_dir|batch_dir> [--out report.json]
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

ANY_NAMING_RE = re.compile(
    r"any (sheet|tab|worksheet)s?[ /a-z]{0,20}(nam|layout|order|count)|"
    r"(sheet|tab)[ /a-z]{0,12}nam\w*[^.]{0,30}\b(is|are) (fine|up to|flexible|your choice)|"
    r"any names?/layout|(sheet|tab) names? (are|is) (not graded|flexible|up to you)",
    re.I,
)
FORMULA_ASK_RE = re.compile(
    r"\bformulas?\b|\blive (calc\w*|link\w*|cell\w*|model)|\blinked\b|\bcalculat\w*|\bcomput(e|ed|es|ing|ation)\b|"
    r"\brecalc\w*|\bCOUNTIFS?\b|\bSUMIFS?\b|\bSUMPRODUCT\b|\bVLOOKUP\b|\bXLOOKUP\b|=SUM\(",
    re.I,
)
CRIT_FORMULA_RE = re.compile(r"formula|\blive\b|recalc|computed|calculated", re.I)
RECALC_RE = re.compile(r"soffice|libreoffice|recalc", re.I)
FORMULA_PATHS = {
    "$.n_formula_with_refs": "n_formula_with_refs",
    "$.n_formula": "n_formula",
    "$.n_cross_sheet_formula": "n_cross_sheet_formula",
    "$.formula_share_all": "formula_share_all",
    "$.formula_share": "formula_share_all",
    "$.formula_share_min": "formula_share_all",
}
FORMULA_COMMANDS = {
    "formula_summary",
    "formula_authenticity_scan",
    "formula_authenticity_check",
}
FLOOR_CMP = {"greater_than_equal", "greater_than", "gte", "gt", ">=", ">"}
# the vendored library's own reference regex (sources/xlsx.py _CELL_REF_RE), so golden counts match what the verifier counts
CELL_REF_RE = re.compile(r"(\$?[A-Za-z]{1,3}\$?\d{1,7})|!|\[[^\]]+\]")
NUMERIC_TEXT_RE = re.compile(r"^[\s$€£(]*[-+]?\d[\d,]*(\.\d+)?\s*[%x)]*\s*$")


# ----------------------------------------------------------------------------- package model


def resolve_package_dir(task_dir: Path) -> Path | None:
    if (task_dir / "tests").is_dir():
        return task_dir
    if (task_dir / "task_files").is_dir():
        return task_dir / "task_files"
    revs = sorted(task_dir.glob("rev_*/task_files")) or sorted(
        task_dir.glob("*/rev_*/task_files")
    )
    return revs[-1] if revs else None


def read_instruction(package: Path) -> str:
    out = ""
    for p in (package / "instruction.md", package / "tests" / "task_prompt.txt"):
        if p.is_file():
            out += p.read_text(encoding="utf-8", errors="replace") + "\n"
    return out


def _rel(package: Path, p: Path) -> str:
    try:
        return p.relative_to(package).as_posix()
    except ValueError:
        return str(p)


def golden_workbook(package: Path, deliverable: str) -> Path | None:
    name = Path(deliverable).name
    sol = package / "solution"
    if not sol.is_dir() or not name:
        return None
    hits = sorted(sol.rglob(name))
    return hits[0] if hits else None


def workbook_formula_stats(path: Path) -> dict:
    """Same counting rules as the vendored verifier_lib summarize_formulas / scan_formula_authenticity."""
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=False, read_only=False)
    n_formula = n_refs = n_cross = hard = 0
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    v = cell.value
                    if cell.data_type == "f":
                        n_formula += 1
                        text = (
                            v
                            if isinstance(v, str)
                            else (getattr(v, "text", None) or "")
                        )
                        if CELL_REF_RE.search(text):
                            n_refs += 1
                        if isinstance(v, str) and "!" in v:
                            n_cross += 1
                    elif isinstance(v, bool):
                        continue
                    elif isinstance(v, (int, float)) or (
                        isinstance(v, str)
                        and NUMERIC_TEXT_RE.match(v)
                        and any(ch.isdigit() for ch in v)
                    ):
                        hard += 1
    finally:
        wb.close()
    checked = n_refs + hard
    return {
        "n_formula": n_formula,
        "n_formula_with_refs": n_refs,
        "n_cross_sheet_formula": n_cross,
        "formula_share_all": round(n_refs / checked, 4) if checked else 0.0,
    }


def native_verifiers(package: Path) -> list[dict]:
    vj = package / "tests" / "verifier.json"
    if not vj.is_file():
        return []
    try:
        data = json.loads(vj.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return []
    out = []
    for v in data.get("verifiers", []) or []:
        a = v.get("assertion") or {}
        f = (
            ((v.get("source") or {}).get("file") or {})
            if isinstance(v.get("source"), dict)
            else {}
        )
        det = a.get("deterministic") or {}
        out.append(
            {
                "name": str(v.get("name", "?")),
                "type": a.get("type"),
                "stype": f.get("type"),
                "command": f.get("command"),
                "args": f.get("arguments") or {},
                "det_path": det.get("path"),
                "comparison": str(det.get("comparison", "")).lower(),
                "expected": a.get("expected"),
            }
        )
    return out


def _res(status: str, evidence: list[str]) -> dict:
    return {"status": status, "evidence": evidence}


def is_formula_gate(v: dict) -> bool:
    return v["type"] == "deterministic" and (
        v["command"] in FORMULA_COMMANDS or v["det_path"] in FORMULA_PATHS
    )


# ----------------------------------------------------------------------------- rules


def g1_layout_pin(package: Path, verifiers: list[dict], instruction: str) -> dict:
    m = ANY_NAMING_RE.search(instruction)
    if not m:
        return _res("pass", ["instruction does not allow free tab naming"])
    ev = []
    for v in verifiers:
        if v["type"] != "deterministic":
            continue
        sheet = v["args"].get("sheet")
        rng = v["args"].get("cell_range")
        if not (sheet or rng):
            continue
        if sheet and str(sheet) in instruction:
            continue  # the instruction itself prescribes this tab name
        ev.append(
            f"tests/verifier.json {v['name']}: {v['command']} pins sheet={sheet!r} cell_range={rng!r} while instruction.md says "
            f"'{m.group(0)}'. Fix: use a layout-independent command (formula_summary / formula_authenticity_scan) or a rubric "
            f"criterion, or state the required tab name and range in the instruction"
        )
    return _res(
        "warn" if ev else "pass",
        ev or ["no sheet/cell_range pin on a deterministic verifier"],
    )


def h1_formula_floor(package: Path, verifiers: list[dict]) -> dict:
    ev = []
    seen = 0
    for v in verifiers:
        if (
            v["type"] != "deterministic"
            or v["det_path"] not in FORMULA_PATHS
            or v["comparison"] not in FLOOR_CMP
        ):
            continue
        try:
            floor = float(v["expected"])
        except (TypeError, ValueError):
            continue
        deliv = str(v["args"].get("path", ""))
        gold = golden_workbook(package, deliv)
        if gold is None:
            continue
        seen += 1
        try:
            stats = workbook_formula_stats(gold)
        except Exception as exc:  # noqa: BLE001
            ev.append(
                f"tests/verifier.json {v['name']}: could not read golden {_rel(package, gold)}: {exc}"
            )
            continue
        g = stats[FORMULA_PATHS[v["det_path"]]]
        strict = v["comparison"] in {"greater_than", "gt", ">"}
        if g < floor or (strict and g <= floor):
            verdict = "the golden itself fails this floor"
        elif g == floor:
            verdict = "zero headroom: a compliant layout with one fewer formula fails"
        else:
            continue
        ev.append(
            f"tests/verifier.json {v['name']}: {v['det_path']} {v['comparison']} {v['expected']} but golden "
            f"{_rel(package, gold)} has {g} ({verdict}). Fix: set the floor to what the instruction requires "
            f"(count the formulas the prompt asks for), below the golden's own count"
        )
    if not seen:
        return _res("skip", ["no formula floor with a matching golden workbook"])
    return _res(
        "warn" if ev else "pass",
        ev or [f"{seen} formula floor(s) below the golden's own value"],
    )


def h2_formula_unasked(package: Path, verifiers: list[dict], instruction: str) -> dict:
    gates = [v for v in verifiers if is_formula_gate(v)]
    if not gates:
        return _res("skip", ["no formula-count/share verifier"])
    if FORMULA_ASK_RE.search(instruction):
        return _res(
            "pass",
            [
                f"instruction asks for formulas/calculation ('{FORMULA_ASK_RE.search(instruction).group(0)}')"
            ],
        )
    ev = [
        f"tests/verifier.json {v['name']}: {v['command']} {v['det_path'] or ''} {v['comparison']} {v['expected']} "
        f"on {v['args'].get('path', '?')}, but instruction.md never asks for formulas, live links, or calculated cells. "
        f"Fix: remove the formula gate, or add the formula requirement to the instruction"
        for v in gates
    ]
    return _res("warn", ev)


def i1_rewardkit_no_recalc(package: Path) -> dict:
    import tomllib

    tomls = sorted((package / "tests").glob("*/judge.toml"))
    if not tomls:
        return _res("skip", ["not a rewardkit (judge.toml) task"])
    env_names = (
        {p.name for p in (package / "environment").rglob("*") if p.is_file()}
        if (package / "environment").is_dir()
        else set()
    )
    hits = []
    for jt in tomls:
        try:
            d = tomllib.loads(jt.read_text(encoding="utf-8", errors="replace"))
        except Exception:  # noqa: BLE001
            continue
        files = [str(x) for x in ((d.get("judge") or {}).get("files") or [])]
        xl = [
            f
            for f in files
            if f.lower().endswith((".xlsx", ".xlsm")) and Path(f).name not in env_names
        ]
        if not xl:
            continue
        crits = [
            str(c.get("name", "?"))
            for c in (d.get("criterion") or d.get("criteria") or [])
            if CRIT_FORMULA_RE.search(str(c.get("description", "")))
        ]
        if crits:
            hits.append((jt, xl, crits))
    if not hits:
        return _res(
            "pass", ["no judge group grades formulas in a deliverable workbook"]
        )
    blob = ""
    for p in (package / "tests" / "test.sh", package / "environment" / "Dockerfile"):
        if p.is_file():
            blob += p.read_text(encoding="utf-8", errors="replace")
    if RECALC_RE.search(blob):
        return _res(
            "pass",
            [
                f"{len(hits)} formula-grading group(s); recalc step present in test.sh/Dockerfile"
            ],
        )
    ev = [
        f"{_rel(package, jt)}: files={xl} criteria {crits[:6]} grade formulas/computed values, but tests/test.sh and "
        f"environment/Dockerfile have no soffice/libreoffice recalc step, so an agent's uncached formulas reach the judge "
        f"as blank cells. Fix: add a headless LibreOffice recalc of output/*.xlsx in test.sh before rewardkit runs "
        f"(and install libreoffice-calc in the Dockerfile)"
        for jt, xl, crits in hits
    ]
    return _res("warn", ev)


def check_task(task_dir: Path) -> dict:
    package = resolve_package_dir(task_dir)
    if package is None:
        return {"status": "skip", "checks": {}}
    instruction = read_instruction(package)
    verifiers = native_verifiers(package)
    checks = {}
    if verifiers:
        checks["G1_layout_pin"] = g1_layout_pin(package, verifiers, instruction)
        checks["H1_formula_floor"] = h1_formula_floor(package, verifiers)
        checks["H2_formula_unasked"] = h2_formula_unasked(
            package, verifiers, instruction
        )
    checks["I1_rk_no_recalc"] = i1_rewardkit_no_recalc(package)
    order = {"fail": 3, "warn": 2, "pass": 1, "skip": 0}
    worst = max(
        (c["status"] for c in checks.values()), key=lambda s: order[s], default="skip"
    )
    return {
        "status": worst,
        "advisory": True,
        "task": task_dir.name,
        "package": str(package),
        "checks": checks,
    }


def discover(root: Path) -> list[Path]:
    root = root.resolve()
    if (
        (root / "tests").is_dir()
        or (root / "task_files").is_dir()
        or list(root.glob("rev_*/task_files"))
    ):
        return [root]
    return sorted(p for p in root.iterdir() if p.is_dir())


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    ap.add_argument("--root", type=Path, required=True)
    ap.add_argument("--out", type=Path)
    ap.add_argument("--task", action="append", dest="task_ids")
    args = ap.parse_args(argv)
    dirs = discover(args.root)
    if args.task_ids:
        want = set(args.task_ids)
        dirs = [d for d in dirs if d.name in want]
    rows = {}
    for d in dirs:
        r = check_task(d)
        rows[d.name] = r
        warns = [k for k, v in r.get("checks", {}).items() if v["status"] == "warn"]
        print(f"{d.name:24} {r['status']:5} warn={warns}")
    if args.out:
        args.out.parent.mkdir(parents=True, exist_ok=True)
        args.out.write_text(
            json.dumps(
                {"generated": datetime.now(timezone.utc).isoformat(), "tasks": rows},
                indent=2,
            ),
            encoding="utf-8",
        )
        print(f"wrote {args.out}")
    return 0  # advisory: never blocks


if __name__ == "__main__":
    sys.exit(main())
