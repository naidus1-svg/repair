"""Deliverable integrity check: confirm every file under a task's golden solution/ directory
actually opens/parses without corruption, based on its extension.

Why this exists: a build/conversion/upload step (an xlsm->xlsx shim, a zip round-trip, a Drive
push, a partial write on a crash) can leave a golden deliverable that looks present on disk --
right name, plausible size -- but is actually truncated or malformed, and a plain existence check
(does this path exist?) never catches that. This script opens each file with the parser that
actually matters for its format and reports whether that parser accepts it.

Provenance: new script, generalizing an existing narrow pattern in this repo (the tracker
builder's own `zipfile.ZipFile(path).testzip()` step, used only to verify the tracker workbook it
had just built) into a per-task, per-file, per-format check over every deliverable a task ships,
not just one known-format workbook. It expects the same generic Harbor task package shape as the
other scripts in this repo: for each task, a `<task_dir>/solution/` folder holding the golden
deliverable file(s), in whatever formats a given task happens to require. If your package shape
differs, `find_deliverables` is the one place to adapt.

Status: v1, generalizes an existing pattern (zipfile.testzip on a built workbook) to every
deliverable file type, but the .pdf structural check is a basic well-formedness heuristic, not a
full validator.

Usage:
    python deliverable_integrity_check.py <TASK-ID> [<TASK-ID> ...] --root <path-to-task-dirs> --out report.json
    python deliverable_integrity_check.py --all --root <path-to-task-dirs> --out report.json
    python deliverable_integrity_check.py --all --root <path> --id-prefix MYBENCH- --out report.json

What it does, in order:

1. For a given task id (or --all tasks under --root), finds every file under that task's
   solution/ directory (recursive glob -- every file, not just spreadsheets, since a task's
   golden deliverable can be a workbook, a Word doc, a PDF, or something else entirely).
2. Dispatches each file to a check based on its extension (case-insensitive):
     - .xlsx / .xlsm : opened with `openpyxl.load_workbook`, then every worksheet's every cell
                        is iterated (forces the parse to actually touch the whole file rather
                        than trusting that `load_workbook` returning without raising means the
                        content is intact). Any exception during load or iteration is a fail.
     - .docx          : a .docx IS a zip archive, so it is opened with `zipfile.ZipFile` and
                        checked two ways: `testzip()` (validates every member's CRC and returns
                        the name of the first corrupt one, or None if all are fine -- the exact
                        pattern this script generalizes from the tracker builder), and a check
                        that `word/document.xml` is actually present in the archive. A zip can
                        pass testzip() cleanly and still not be a Word document at all (wrong
                        Office format saved under a .docx extension, or a zip missing that one
                        member) -- testzip() alone would call that a pass, so both checks are
                        required.
     - .pdf           : no parsing library is added for this (a full PDF parser is a heavy
                        dependency for a v1 "does this look well-formed" check). Instead a
                        lightweight pure-stdlib structural check: the file must start with the
                        `%PDF-` header, contain at least one `endobj` marker (i.e. it has SOME
                        object structure, not just a header slapped onto garbage), and carry a
                        `startxref` pointer plus a trailing `%%EOF` -- the standard PDF trailer
                        machinery -- tolerating either a classical `trailer` dictionary or a
                        cross-reference stream (`/Type /XRef`), since a modern or linearized PDF
                        often uses the latter with no `trailer` keyword anywhere in the file.
                        This is disclosed below as a heuristic, not a real validator: it cannot
                        catch, for example, a structurally valid-looking PDF whose content
                        streams are internally corrupt.
     - anything else  : reported as `"not checked, unknown format"` with `checked: false` --
                        never silently dropped from the file list, just not asserted either way.
3. Writes one JSON result per task (status: pass|fail, per-file findings, summary counts) plus an
   aggregate report across every task given on the command line (or discovered via --all). A task
   fails if and only if at least one checked file failed its check; unchecked (unknown-format)
   files never fail a task on their own.

Known limits, disclosed rather than silently dropped:
    - The .pdf check is structural well-formedness only, not full validation. A PDF can pass this
      check and still have a corrupt content stream, a broken font, or content a renderer
      chokes on. It catches truncation, a missing/garbled trailer, and non-PDF files wearing a
      .pdf extension -- not everything a real PDF parser would.
    - The .docx check confirms the archive is intact and contains word/document.xml. It does not
      parse that XML, so a document.xml that is present but itself malformed (e.g. truncated
      mid-tag) is a case this check does not have a stdlib-only path to catch beyond what
      zipfile.testzip() already validates about the compressed bytes; a genuinely truncated XML
      member usually also fails testzip() or fails to be read back, both of which this check does
      catch.
    - .xls (legacy binary Excel), .pptx, .csv, and every other extension besides the four handled
      above are reported as "not checked, unknown format" in this v1 -- not asserted to be fine,
      just not covered yet.
    - This script analyzes the CURRENT on-disk state of solution/ under --root. It says nothing
      about whether a task's deliverable was broken in the past and has since been fixed, and
      nothing about environment/initial_workspace/ or any file outside solution/.
"""

# Copied into turing-context on 2026-09-14 from harbor-qc-checks/deliverable_integrity_check.py.
# See tools/svc_qc_tracker_build.py CHECK_META["D"]["subchecks"]["D7"]. No logic changed.

import argparse
import json
import re
import sys
import zipfile
from datetime import datetime, timezone
from pathlib import Path

# Windows console defaults to cp1252, which cannot encode every character that can show up in a
# library's own exception text. Force UTF-8 with replacement so a printable-but-non-cp1252
# character never crashes the whole sweep partway through a batch.
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.stderr.reconfigure(encoding="utf-8", errors="replace")

try:
    import openpyxl
except ImportError:  # pragma: no cover
    import subprocess

    subprocess.run(
        [sys.executable, "-m", "pip", "install", "openpyxl", "--quiet"], check=True
    )
    import openpyxl

SPREADSHEET_SUFFIXES = (".xlsx", ".xlsm")
XREF_STREAM_RE = re.compile(rb"/Type\s*/XRef")


# ---------------------------------------------------------------------------
# task/file discovery
# ---------------------------------------------------------------------------


def find_task_dirs(root, task_ids=None, id_prefix=""):
    if task_ids:
        dirs = []
        for tid in task_ids:
            p = root / tid
            if not p.is_dir():
                print(f"WARNING: task dir not found, skipping: {p}", file=sys.stderr)
                continue
            dirs.append(p)
        return dirs
    return sorted(
        p for p in root.iterdir() if p.is_dir() and p.name.startswith(id_prefix)
    )


def find_deliverables(task_dir):
    """Every file under <task_dir>/solution/, recursive -- not just spreadsheets, since a
    task's golden deliverable can be any file type."""
    sol = task_dir / "solution"
    if not sol.is_dir():
        return []
    return sorted(p for p in sol.rglob("*") if p.is_file())


# ---------------------------------------------------------------------------
# per-format checks
# ---------------------------------------------------------------------------


def check_spreadsheet(path):
    """.xlsx / .xlsm: open with openpyxl and force a full read of every cell on every sheet, so
    a corruption that only surfaces once the parser actually touches the cell data (as opposed
    to just the archive/workbook-level structure `load_workbook` validates on its own) is not
    missed."""
    try:
        wb = openpyxl.load_workbook(str(path), data_only=False)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    _ = cell.value
        wb.close()
    except Exception as e:
        return {
            "ok": False,
            "detail": f"openpyxl could not open/parse this workbook: {e}",
        }
    return {
        "ok": True,
        "detail": "opened with openpyxl and every cell on every sheet was read cleanly",
    }


def check_docx(path):
    """.docx IS a zip archive. testzip() validates every member's CRC and returns the name of
    the first corrupt member (None if all are fine) -- the exact pattern this script generalizes
    from the tracker builder's own verification of a workbook it had just built. That alone is
    not sufficient: a zip can pass testzip() and still not be a Word document (wrong Office
    format under a .docx extension, or missing the one member that makes it a Word document), so
    word/document.xml presence is checked too."""
    try:
        with zipfile.ZipFile(path) as zf:
            bad_member = zf.testzip()
            if bad_member is not None:
                return {
                    "ok": False,
                    "detail": f"zip CRC check failed on member: {bad_member}",
                }
            names = zf.namelist()
            if "word/document.xml" not in names:
                return {
                    "ok": False,
                    "detail": "valid zip archive, CRCs pass, but no word/document.xml member -- "
                    "not a Word document (wrong Office format, or a corrupt/incomplete docx)",
                }
    except zipfile.BadZipFile as e:
        return {"ok": False, "detail": f"not a valid zip archive: {e}"}
    except OSError as e:
        return {"ok": False, "detail": f"could not read file: {e}"}
    return {
        "ok": True,
        "detail": "valid zip archive, every member's CRC checks out, word/document.xml present",
    }


def check_pdf(path):
    """Basic structural well-formedness check for a PDF, pure stdlib, no parsing library added.
    Confirms: the file starts with the %PDF- header, contains at least one endobj marker (some
    object structure exists, not just a header slapped onto garbage), and carries a startxref
    pointer plus a trailing %%EOF -- the standard PDF trailer machinery -- tolerating either a
    classical `trailer` dictionary or a cross-reference stream (/Type /XRef), since a modern or
    linearized PDF often uses the latter with no "trailer" keyword anywhere in the file. This is
    a heuristic, not a full validator -- see the module docstring's Known limits."""
    try:
        data = path.read_bytes()
    except OSError as e:
        return {"ok": False, "detail": f"could not read file: {e}"}

    if not data:
        return {"ok": False, "detail": "file is empty (0 bytes)"}

    if not data.startswith(b"%PDF-"):
        return {"ok": False, "detail": "does not start with the %PDF- header"}

    if b"endobj" not in data:
        return {
            "ok": False,
            "detail": "no 'endobj' marker found anywhere in the file -- no object structure",
        }

    tail = data[-2048:] if len(data) > 2048 else data
    if b"startxref" not in data:
        return {
            "ok": False,
            "detail": "no 'startxref' marker found -- missing cross-reference pointer",
        }
    if b"%%EOF" not in tail:
        return {
            "ok": False,
            "detail": "no trailing '%%EOF' marker found near the end of the file",
        }
    has_trailer_dict = b"trailer" in data
    has_xref_stream = XREF_STREAM_RE.search(data) is not None
    if not (has_trailer_dict or has_xref_stream):
        return {
            "ok": False,
            "detail": "no classical 'trailer' dictionary and no cross-reference stream "
            "(/Type /XRef) found -- neither trailer mechanism present",
        }

    return {
        "ok": True,
        "detail": "starts with %PDF-, has endobj/startxref/%%EOF and a trailer mechanism "
        "(classical trailer dict or xref stream)",
    }


EXT_CHECKERS = {
    ".xlsx": ("spreadsheet", check_spreadsheet),
    ".xlsm": ("spreadsheet", check_spreadsheet),
    ".docx": ("docx", check_docx),
    ".pdf": ("pdf", check_pdf),
}


def check_file(path):
    suffix = path.suffix.lower()
    if suffix in EXT_CHECKERS:
        kind, fn = EXT_CHECKERS[suffix]
        result = fn(path)
        result["kind"] = kind
        result["checked"] = True
        return result
    return {
        "kind": "unknown",
        "checked": False,
        "ok": None,
        "detail": "not checked, unknown format",
    }


# ---------------------------------------------------------------------------
# per-task analysis
# ---------------------------------------------------------------------------


def analyze_task(task_dir):
    task_id = task_dir.name
    sol_exists = (task_dir / "solution").is_dir()
    files = find_deliverables(task_dir)

    result = {
        "task_id": task_id,
        "solution_dir_exists": sol_exists,
        "files": [],
        "errors": [],
    }

    if not files:
        result["status"] = "pass"
        result["summary"] = {
            "n_files": 0,
            "n_checked": 0,
            "n_not_checked": 0,
            "n_corrupt": 0,
        }
        result["errors"].append(
            "no files found under solution/ -- nothing for this check to verify"
            + ("" if sol_exists else " (solution/ directory itself does not exist)")
        )
        return result

    n_corrupt = 0
    n_not_checked = 0
    for path in files:
        rel = str(path.relative_to(task_dir)).replace("\\", "/")
        try:
            r = check_file(path)
        except Exception as e:  # never let one file's crash kill the whole task
            r = {
                "kind": "error",
                "checked": True,
                "ok": False,
                "detail": f"unhandled exception while checking this file: {e}",
            }
        entry = {
            "path": rel,
            "extension": path.suffix.lower() or "(none)",
            "kind": r["kind"],
            "checked": r["checked"],
            "ok": r["ok"],
            "detail": r["detail"],
        }
        result["files"].append(entry)
        if r["checked"] and r["ok"] is False:
            n_corrupt += 1
        if not r["checked"]:
            n_not_checked += 1

    result["summary"] = {
        "n_files": len(files),
        "n_checked": len(files) - n_not_checked,
        "n_not_checked": n_not_checked,
        "n_corrupt": n_corrupt,
    }
    result["status"] = "fail" if n_corrupt else "pass"
    return result


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------


def main():
    ap = argparse.ArgumentParser(
        description="Confirm every file under a task's golden solution/ directory actually "
        "opens/parses without corruption, based on its extension (.xlsx/.xlsm via openpyxl, "
        ".docx via zipfile, .pdf via a basic structural check; anything else reported as not "
        "checked rather than silently skipped)."
    )
    ap.add_argument(
        "task_ids",
        nargs="*",
        help="one or more task ids (directory names under --root)",
    )
    ap.add_argument(
        "--all", action="store_true", help="run every task directory found under --root"
    )
    ap.add_argument(
        "--root",
        type=Path,
        default=None,
        required=False,
        help="root directory containing one subdirectory per task (REQUIRED -- no default, "
        "this is a generalized script that must not silently assume a project-specific path)",
    )
    ap.add_argument(
        "--id-prefix",
        default="",
        help="when using --all, only treat subdirectories starting with this prefix as tasks "
        "(default: empty, i.e. every subdirectory is treated as a task)",
    )
    ap.add_argument(
        "--out",
        type=Path,
        default=None,
        help="write the aggregate JSON report here (default: print summary only)",
    )
    args = ap.parse_args()

    if args.root is None:
        ap.error(
            "--root is required (point it at the directory containing your task folders)"
        )

    if not args.all and not args.task_ids:
        ap.error("pass one or more task ids, or --all")

    task_dirs = find_task_dirs(
        args.root, None if args.all else args.task_ids, id_prefix=args.id_prefix
    )
    if not task_dirs:
        print("No task directories found.", file=sys.stderr)
        sys.exit(1)

    results = []
    for i, task_dir in enumerate(task_dirs, 1):
        print(f"[{i}/{len(task_dirs)}] {task_dir.name} ...", file=sys.stderr)
        try:
            r = analyze_task(task_dir)
        except Exception as e:  # never let one task's crash kill the whole batch
            r = {
                "task_id": task_dir.name,
                "status": "fail",
                "errors": [f"unhandled exception analyzing this task: {e}"],
                "files": [],
                "summary": {
                    "n_files": 0,
                    "n_checked": 0,
                    "n_not_checked": 0,
                    "n_corrupt": 0,
                },
            }
        results.append(r)
        print(
            f"    -> {r['status']}  "
            f"({r.get('summary', {}).get('n_corrupt', 0)} corrupt, "
            f"{r.get('summary', {}).get('n_not_checked', 0)} unchecked, "
            f"{len(r.get('errors', []))} errors)",
            file=sys.stderr,
        )

    n_pass = sum(1 for r in results if r["status"] == "pass")
    n_fail = sum(1 for r in results if r["status"] == "fail")
    aggregate = {
        "generated": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "root": str(args.root),
        "n_tasks": len(results),
        "n_pass": n_pass,
        "n_fail": n_fail,
        "results": results,
    }

    print(f"\n{n_pass} pass / {n_fail} fail out of {len(results)} task(s)")
    for r in results:
        corrupt_paths = [
            f["path"]
            for f in r.get("files", [])
            if f.get("checked") and f.get("ok") is False
        ]
        suffix = f"  [corrupt: {', '.join(corrupt_paths)}]" if corrupt_paths else ""
        print(f"  {r['task_id']}: {r['status']}{suffix}")
        for e in r.get("errors", []):
            print(f"      ! {e}")

    if args.out:
        args.out.parent.mkdir(parents=True, exist_ok=True)
        with open(args.out, "w", encoding="utf-8") as f:
            json.dump(aggregate, f, indent=2)
        print(f"\nWrote {args.out}")


if __name__ == "__main__":
    main()
