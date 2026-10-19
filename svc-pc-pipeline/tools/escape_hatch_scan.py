#!/usr/bin/env python3
"""
escape_hatch_scan.py - GATE-V2 class M: criterion escape hatches and
L-series authoring labels.

Per task, flags three things (GATE-V2.md 1b class M, 3.3/3.4 "M, escape
hatches", section 4 step 1 `escape_hatch_scan.py`):

  (a) criterion text (in any tests/deliverables*/judge.toml) naming a file
      path that is NOT in that judge group's `files=` list -- a criterion
      that can be satisfied by pointing at material the judge never sees.
  (b) criterion text containing the tokens worker, upstream, companion,
      "prior deliverable", or an authoring-series label matching the regex
      L[0-9]-[0-9][0-9] (e.g. "L1-02-01").
  (c) the same L[0-9]-[0-9][0-9] pattern in golden cell text and xlsx sheet
      titles under solution/files (docx paragraphs/tables and other
      formats are also checked via the shared text extractor).

READ-ONLY: only reads corpus files, never modifies them. Run corpus-wide.

Output: CSV [task, where, id_or_loc, token_or_path] where `where` is one
of criterion | golden | sheet-title.

Usage:
    python3 escape_hatch_scan.py [--corpus-root PATH] [--out PATH] [--task NAME]
"""
from __future__ import annotations

import argparse
import re
import sys
import tomllib
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _qc_common import (
    DEFAULT_CORPUS_ROOT,
    DEFAULT_REPORTS_DIR,
    extract_document_text,
    extract_filename_tokens,
    iter_deliverable_groups,
    list_tasks,
    write_csv,
)

L_SERIES_RE = re.compile(r"\bL[0-9]-[0-9][0-9](?:-[0-9]+)?\b")
APP_PATH_RE = re.compile(r"/app/[A-Za-z0-9_./\-]+\.[A-Za-z0-9]{2,5}")
TOKEN_WORDS = ("worker", "upstream", "companion")
PHRASE_TOKENS = ("prior deliverable",)


def _group_name(judge_path: Path) -> str:
    return judge_path.parent.name  # "deliverables" or "deliverables_NN"


def scan_criteria(task_name: str, judge_path: Path) -> list[tuple[str, str, str, str]]:
    rows: list[tuple[str, str, str, str]] = []
    try:
        data = tomllib.loads(judge_path.read_text(encoding="utf-8", errors="ignore"))
    except Exception:
        return rows

    group = _group_name(judge_path)
    files = data.get("judge", {}).get("files", []) or []
    basenames = {Path(f).name.lower() for f in files}

    for c in data.get("criterion", []) or []:
        cid = c.get("name", "?")
        desc = c.get("description", "") or ""
        loc = f"{group}:{cid}"

        # (a) escape-hatch path/filename not in this group's files=
        candidates = set(APP_PATH_RE.findall(desc)) | extract_filename_tokens(desc)
        for cand in sorted(candidates):
            bn = Path(cand).name.lower()
            if bn and bn not in basenames:
                rows.append((task_name, "criterion", loc, cand))

        # (b) worker/upstream/companion/prior-deliverable/L-series tokens
        low = desc.lower()
        for w in TOKEN_WORDS:
            if re.search(r"\b%s\b" % re.escape(w), low):
                rows.append((task_name, "criterion", loc, w))
        for phrase in PHRASE_TOKENS:
            if phrase in low:
                rows.append((task_name, "criterion", loc, phrase))
        for m in sorted(set(L_SERIES_RE.findall(desc))):
            rows.append((task_name, "criterion", loc, m))

    return rows


def scan_golden(task_name: str, task_dir: Path) -> list[tuple[str, str, str, str]]:
    rows: list[tuple[str, str, str, str]] = []
    sol_dir = task_dir / "solution" / "files"
    if not sol_dir.is_dir():
        return rows

    for f in sorted(sol_dir.iterdir()):
        if not f.is_file():
            continue
        ext = f.suffix.lower()
        rel = str(f.relative_to(task_dir))

        if ext == ".xlsx":
            try:
                import openpyxl

                wb = openpyxl.load_workbook(f, data_only=True, read_only=True)
                for ws in wb.worksheets:
                    for m in sorted(set(L_SERIES_RE.findall(ws.title))):
                        rows.append((task_name, "sheet-title", f"{rel}!{ws.title}", m))
                    for row in ws.iter_rows():
                        for cell in row:
                            v = cell.value
                            if isinstance(v, str) and L_SERIES_RE.search(v):
                                for m in sorted(set(L_SERIES_RE.findall(v))):
                                    rows.append(
                                        (task_name, "golden", f"{rel}!{ws.title}!{cell.coordinate}", m)
                                    )
                wb.close()
            except Exception:
                pass
        elif ext == ".docx":
            try:
                import docx

                d = docx.Document(str(f))
                for i, p in enumerate(d.paragraphs):
                    if L_SERIES_RE.search(p.text):
                        for m in sorted(set(L_SERIES_RE.findall(p.text))):
                            rows.append((task_name, "golden", f"{rel}!para{i}", m))
                for ti, table in enumerate(d.tables):
                    for ri, trow in enumerate(table.rows):
                        for ci, cell in enumerate(trow.cells):
                            if L_SERIES_RE.search(cell.text):
                                for m in sorted(set(L_SERIES_RE.findall(cell.text))):
                                    rows.append(
                                        (task_name, "golden", f"{rel}!table{ti}r{ri}c{ci}", m)
                                    )
            except Exception:
                pass
        else:
            text = extract_document_text(f)
            if text and L_SERIES_RE.search(text):
                for m in sorted(set(L_SERIES_RE.findall(text))):
                    rows.append((task_name, "golden", f"{rel}!body", m))
    return rows


def scan_task(task_dir: Path) -> list[tuple[str, str, str, str]]:
    rows: list[tuple[str, str, str, str]] = []
    for judge_path in iter_deliverable_groups(task_dir):
        jt = judge_path / "judge.toml"
        if jt.is_file():
            rows.extend(scan_criteria(task_dir.name, jt))
    rows.extend(scan_golden(task_dir.name, task_dir))
    return rows


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--corpus-root", type=Path, default=DEFAULT_CORPUS_ROOT)
    ap.add_argument("--out", type=Path, default=DEFAULT_REPORTS_DIR / "escape_hatch.csv")
    ap.add_argument("--task", type=str, default=None)
    ap.add_argument("--verdict-out", type=Path, default=None, help="write task_id,escape_hatch_scan pass|fail summary CSV")
    args = ap.parse_args()

    tasks = list_tasks(args.corpus_root)
    if args.task:
        tasks = [t for t in tasks if t.name == args.task]
    if not tasks:
        print(f"ERROR: no matching task directories under {args.corpus_root}", file=sys.stderr)
        return 1

    all_rows: list[tuple[str, str, str, str]] = []
    errors: list[str] = []
    for i, task_dir in enumerate(tasks, 1):
        try:
            all_rows.extend(scan_task(task_dir))
        except Exception as exc:
            errors.append(f"{task_dir.name}: {exc}")
        if i % 25 == 0:
            print(f"  ...{i}/{len(tasks)} tasks scanned", file=sys.stderr)

    n = write_csv(args.out, ["task", "where", "id_or_loc", "token_or_path"], all_rows)
    tasks_flagged = {r[0] for r in all_rows}
    by_where = {}
    for r in all_rows:
        by_where[r[1]] = by_where.get(r[1], 0) + 1

    print(f"escape_hatch_scan: {n} flags across {len(tasks_flagged)} of {len(tasks)} tasks, written to {args.out}")
    for w, c in sorted(by_where.items()):
        print(f"  {w}: {c}")
    if errors:
        print(f"  {len(errors)} tasks raised errors during scan:", file=sys.stderr)
        for e in errors:
            print(f"    {e}", file=sys.stderr)

    if args.verdict_out:
        verdict_rows = sorted(
            (t.name, "fail" if t.name in tasks_flagged else "pass") for t in tasks
        )
        vn = write_csv(args.verdict_out, ["task_id", "escape_hatch_scan"], verdict_rows)
        print(f"escape_hatch_scan: verdict CSV written, {vn} rows, to {args.verdict_out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
