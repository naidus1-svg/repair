#!/usr/bin/env python3
"""Build the SVC auto-fix summary workbook from a scan + apply run's output.

Reuses the tab pattern from tools/svc_fp_screen_sheet.py: a narrative Summary
tab first, then per-disposition tabs a reader can act on directly.

    python tools/svc_autofix_report.py \\
        --scan ekw/svc-autofix-scan-2026-09-12.json \\
        --apply ekw/svc-autofix-apply-2026-09-12.json \\
        --out ekw/svc-autofix-report-2026-09-12.xlsx
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPO_ROOT = Path(__file__).resolve().parent.parent
DEFAULT_SCAN = REPO_ROOT / "ekw" / "svc-autofix-scan-2026-09-12.json"
DEFAULT_APPLY = REPO_ROOT / "ekw" / "svc-autofix-apply-2026-09-12.json"
DEFAULT_OUT = REPO_ROOT / "ekw" / "svc-autofix-report-2026-09-12.xlsx"

FINDING_COLUMNS = [
    "QC Issues Task ID",
    "Defect Class",
    "Auto-Verdict",
    "Fix Title",
    "Issue Description",
    "Evidence Reference (path)",
]
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="343A40")
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
BUCKET_FILL = {
    "mechanical": PatternFill("solid", fgColor="D4EDDA"),
    "content_review": PatternFill("solid", fgColor="FFF3CD"),
    "no_action": PatternFill("solid", fgColor="E2E3E5"),
    "needs_human": PatternFill("solid", fgColor="F8D7DA"),
}


def style_header(ws, labels: list[str]) -> None:
    ws.append(labels)
    for col in range(1, len(labels) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = WRAP_TOP
    ws.freeze_panes = "A2"


def build_summary(wb, scan: dict, apply_data: "dict | None") -> None:
    ws = wb.active
    ws.title = "Summary"
    counts = scan.get("counts", {})
    rows = [
        ("SVC auto-fix: run summary", ""),
        ("Generated", scan.get("generated", "")),
        ("", ""),
        ("Tasks scanned", scan.get("task_count", 0)),
        ("Sample requested", scan.get("sample_requested") or "(full run)"),
        (
            "Task selection source",
            (
                "rollup JSON, lowest confirmed_defect_count first"
                if scan.get("rollup_used")
                else "xlsx files directly (rollup not found), first N found"
            ),
        ),
        ("", ""),
        ("Total findings", scan.get("total_findings", 0)),
        ("Mechanical (auto-applied, no review gate)", counts.get("mechanical", 0)),
        (
            "Content review (diff artifact, awaiting Jafar's QC pass)",
            counts.get("content_review", 0),
        ),
        ("No action (already-classified false positive)", counts.get("no_action", 0)),
        (
            "Needs human (false_positive confirmation pending, or genuinely unresolved)",
            counts.get("needs_human", 0),
        ),
        ("", ""),
    ]
    if apply_data:
        mech = apply_data.get("mechanical", {})
        content = apply_data.get("content_review", {})
        rows += [
            ("Apply run generated", apply_data.get("generated", "")),
            ("Mechanical findings applied for real", mech.get("applied", 0)),
            (
                "Mechanical findings skipped (no local task package found)",
                mech.get("skipped_no_local_package", 0),
            ),
            (
                "Mechanical findings with no defect confirmed",
                mech.get("no_defect_confirmed", 0),
            ),
            ("Content diff artifacts written", content.get("total", 0)),
            ("Diff artifact directory", content.get("diff_dir", "")),
            ("", ""),
        ]
    rows += [
        (
            "Hard guard",
            "Every write to a solver-facing path (instruction.md, or anything under an "
            "environment/ directory) is scanned in memory against the task's own "
            "tests/verifier.json expected-value and tolerance strings before it lands; "
            "any overlap refuses the write instead of applying it. See tools/svc_autofix_apply.py.",
        ),
        (
            "Methodology",
            "Defect Class values bucketed per the tp-fp-classification taxonomy "
            "(C:\\ekwqc\\repo\\automation\\tp-fp-classification\\SKILL.md): "
            "format_convention/pipeline_packaging -> mechanical, "
            "content_substance/design_intent -> content_review, "
            "dataos_format_artifact/agent_capability -> no_action, "
            "everything else -> needs_human.",
        ),
    ]
    for r, (label, value) in enumerate(rows, start=1):
        lc = ws.cell(row=r, column=1, value=label)
        vc = ws.cell(row=r, column=2, value=value)
        lc.alignment = WRAP_TOP
        vc.alignment = WRAP_TOP
        if label and not value and r > 3:
            lc.font = Font(bold=True)
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 90


def finding_row(finding: dict) -> list:
    return [finding.get(col, "") for col in FINDING_COLUMNS]


def build_findings(wb, scan: dict) -> None:
    ws = wb.create_sheet("Findings")
    style_header(ws, FINDING_COLUMNS + ["Bucket"])
    r = 2
    for bucket, findings in scan.get("buckets", {}).items():
        for finding in findings:
            row = finding_row(finding) + [bucket]
            for c, value in enumerate(row, start=1):
                cell = ws.cell(row=r, column=c, value=value)
                cell.alignment = WRAP_TOP
            fill = BUCKET_FILL.get(bucket)
            if fill:
                ws.cell(row=r, column=len(FINDING_COLUMNS) + 1).fill = fill
            r += 1
    widths = [18, 16, 14, 30, 45, 40, 16]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width


def build_fixed(wb, apply_data: "dict | None") -> int:
    ws = wb.create_sheet("Fixed")
    style_header(ws, ["Task", "Recipe", "Detail"])
    n = 0
    if apply_data:
        for entry in apply_data.get("mechanical", {}).get("details", []):
            if entry.get("status") != "applied":
                continue
            for recipe in entry.get("recipes", []):
                if not recipe.get("changed"):
                    continue
                n += 1
                ws.append(
                    [
                        entry.get("task_id", ""),
                        recipe.get("recipe", ""),
                        json.dumps(recipe, default=str),
                    ]
                )
    for i, width in enumerate((18, 24, 90), start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    return n


def build_flagged_for_human(wb, scan: dict, apply_data: "dict | None") -> int:
    ws = wb.create_sheet("Flagged for human")
    style_header(ws, ["Task", "Reason", "Fix Title", "Detail / diff artifact"])
    n = 0
    diff_files = (apply_data or {}).get("content_review", {}).get("diff_files", [])
    diff_by_finding = {}
    for path in diff_files:
        stem = Path(path).stem
        diff_by_finding[stem] = path
    for finding in scan.get("buckets", {}).get("content_review", []):
        task_id = finding.get("QC Issues Task ID", "")
        safe_task = "".join(
            ch if ch.isalnum() or ch == "-" else "-" for ch in task_id.lower()
        )
        row_num = finding.get("_row", 0)
        expected_stem = f"{safe_task}__row{row_num}.diff"
        diff_path = diff_by_finding.get(expected_stem, "(not yet applied)")
        ws.append(
            [
                task_id,
                "content/rubric fix awaiting Jafar's QC review",
                finding.get("Fix Title", ""),
                diff_path,
            ]
        )
        n += 1
    for finding in scan.get("buckets", {}).get("needs_human", []):
        ws.append(
            [
                finding.get("QC Issues Task ID", ""),
                f"Auto-Verdict {finding.get('Auto-Verdict', '')}, Defect Class {finding.get('Defect Class', '') or '(blank)'}",
                finding.get("Fix Title", ""),
                finding.get("Actionable Fix Guidance", ""),
            ]
        )
        n += 1
    for i, width in enumerate((18, 40, 30, 60), start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP_TOP
    return n


def main() -> None:
    p = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    p.add_argument("--scan", type=Path, default=DEFAULT_SCAN)
    p.add_argument("--apply", type=Path, default=DEFAULT_APPLY)
    p.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = p.parse_args()

    scan = json.loads(args.scan.read_text(encoding="utf-8"))
    apply_data = (
        json.loads(args.apply.read_text(encoding="utf-8"))
        if args.apply.exists()
        else None
    )

    wb = openpyxl.Workbook()
    build_summary(wb, scan, apply_data)
    build_findings(wb, scan)
    n_fixed = build_fixed(wb, apply_data)
    n_flagged = build_flagged_for_human(wb, scan, apply_data)

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)

    print(f"Wrote {args.out}")
    print(f"  findings: {scan.get('total_findings', 0)}")
    print(f"  fixed: {n_fixed}")
    print(f"  flagged for human: {n_flagged}")
    if apply_data is None:
        print(
            f"  WARNING: {args.apply} not found, Fixed tab is empty and Summary omits apply-run counts"
        )


if __name__ == "__main__":
    main()
